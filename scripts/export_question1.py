#!/usr/bin/env python3
"""Gate, export and read-back audit Q1; never alter the supplied workbooks."""
import csv
import json
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.data_io import ROOT, digest, load_config, save_json, source_fingerprint, verify_sources
from src.boundary import Environment
from src.properties import Diffusivity
from src.fvm import RadialMesh, ScalarFVM
from src.reconstruction import reconstruct
from validation.bessel_q1 import cylinder_response

FIELDS={"temperature":"温度","moisture":"水分浓度"}
PAPER_TIMES=[100,300,600,900,1200,1500,1800]
PAPER_COLUMNS=[0,5,10,15,20]


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def read_case(label):
    # Decimal multipliers (e.g. 0.9) are part of the label, not a file suffix.
    folder=ROOT/"results/cache/q1"
    metadata_path=folder/(label+".json")
    cache_path=folder/(label+".npz")
    meta=read_json(metadata_path)
    if meta["source_fingerprint"]!=source_fingerprint():
        raise ValueError("Stale source/config fingerprint: "+label)
    if meta["script_sha256"]!=digest(ROOT/"scripts/question1.py"):
        raise ValueError("Stale Q1 driver: "+label)
    if meta["cache_sha256"]!=digest(cache_path):
        raise ValueError("Cache hash mismatch: "+label)
    with np.load(cache_path) as raw:
        cache={k:raw[k] for k in raw.files}
    return cache,meta


def export_tables(cache):
    directory=ROOT/"tables"; directory.mkdir(exist_ok=True)
    header=["时间/s","0 cm","0.5 cm","1 cm","1.5 cm","2 cm"]
    markdown=["# 问题一：论文表1—2", "", "由同一全精度基线生成；温度单位 ℃，含水率为 kg/kg（干基）。", ""]
    latex=["% Generated from q1_final full precision; requires booktabs."]
    for index,(field,chinese) in enumerate(FIELDS.items(),1):
        array=cache[field][np.ix_(PAPER_TIMES,PAPER_COLUMNS)]
        rows=[[str(t)]+[f"{v:.4f}" for v in row] for t,row in zip(PAPER_TIMES,array)]
        with (directory/f"q1_{field}.csv").open("w",encoding="utf-8-sig",newline="") as f:
            writer=csv.writer(f); writer.writerow(header); writer.writerows(rows)
        markdown.extend([f"## 表{index}：{chinese}","","| "+" | ".join(header)+" |","|"+"---|"*len(header)])
        markdown.extend("| "+" | ".join(row)+" |" for row in rows)
        markdown.append("")
        caption="不同时刻的径向温度（℃）" if field=="temperature" else "不同时刻的径向干基含水率（kg/kg）"
        latex.extend([r"\begin{table}[htbp]",r"\centering",r"\caption{"+caption+"}",r"\label{tab:q1_"+field+"}",
                      r"\begin{tabular}{rrrrrr}",r"\toprule",r"时间/s & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm \\",r"\midrule"])
        latex.extend(" & ".join(row)+r" \\" for row in rows)
        latex.extend([r"\bottomrule",r"\end{tabular}",r"\end{table}",""])
    (directory/"q1_tables.md").write_text("\n".join(markdown),encoding="utf-8")
    (directory/"q1_tables.tex").write_text("\n".join(latex),encoding="utf-8")


def main():
    original=verify_sources(); cfg=load_config(); env=Environment(cfg)
    cache,meta=read_case("q1_final")
    temporal=read_json(ROOT/"reports/q1_time_and_sensitivity.json")
    spatial=read_json(ROOT/"reports/q1_mesh_convergence.json")
    kernel=read_json(ROOT/"reports/kernel_checks.json")
    assert temporal["complete"] and len(temporal["sensitivities"])==8
    assert meta["perturbation"]=={} and meta["property_set"]=="q1"
    assert (meta["n"],meta["grading"],meta["rtol"],meta["atol"],meta["max_step_s"])==(1280,2.,1e-10,1e-12,1.)
    assert np.array_equal(cache["time_s"],np.arange(1801.))
    assert np.allclose(cache["radius_m"],np.arange(21)*.001,rtol=0,atol=1e-17)
    target=cfg["diagnostics"]["numerical_absolute_target_for_four_decimal_claim"]
    # Check every run's provenance, not just a summary of comparisons.
    for row in spatial["mesh_series"]:
        read_case(row["label"])
    for row in temporal["comparisons"]:
        read_case(row["label"])
    for row in temporal["sensitivities"]:
        read_case(f"sensitivity_{row['parameter']}_{row['multiplier']:g}")
    final_grids=[r for r in spatial["mesh_series"] if r["grading"]==2 and r["n"] in (640,1280)]
    assert len(final_grids)==2
    for row in final_grids:
        assert all(row[field+"_grid_delta_max"]<target for field in FIELDS)
    for row in temporal["comparisons"]:
        if row["label"] in ("cap_2","cap_0p5","radau","kirchhoff","fine2560"):
            assert all(row["comparison_to_reference"][field]["max_abs"]<target for field in FIELDS)
    for row in temporal["step_cap_series"]:
        for field in FIELDS:
            assert abs(row["coarse_fields"][field]["accepted_max_step_s"]-row["coarse_s"])<1e-9
            assert abs(row["fine_fields"][field]["accepted_max_step_s"]-row["fine_s"])<1e-9
            assert row["coarse_fields"][field]["accepted_steps"]!=row["fine_fields"][field]["accepted_steps"]
    assert kernel["constant_D_water_Bessel_max_error"]<target
    with np.load(ROOT/"results/cache/q1/bessel_reference.npz") as raw:
        bessel=raw["temperature"]
    heat_error=float(np.max(np.abs(cache["temperature"][1:]-bessel[1:])))
    assert heat_error<target
    diagnostics={"heat_bessel_full_export_max_abs_C":heat_error,"fields":{}}
    mesh=RadialMesh(meta["n"],cfg["geometry"]["reference_radius_m"],meta["grading"])
    for field in FIELDS:
        values=cache[field]; cells=cache[field+"_cells"]
        assert values.shape==(1801,21) and np.isfinite(cells).all() and np.isfinite(values).all()
        initial=28. if field=="temperature" else 2.55
        assert np.max(np.abs(cells[0]-initial))==0 and np.max(np.abs(values[0]-initial))==0
        assert np.min(cells)>0 and np.min(values)>0
        normalizer=22. if field=="temperature" else 2.55
        independent=meta["fields"][field]["independent_balance_max_abs"]
        assert independent/normalizer<cfg["diagnostics"]["normalized_balance_target"]
        ambient=env.value(cache["time_s"],field)
        law=Diffusivity(.36) if field=="temperature" else Diffusivity(7e-9,.89)
        transfer=25. if field=="temperature" else 8e-7
        capacity=820.*2600 if field=="temperature" else 1.
        model=ScalarFVM(mesh,law,transfer,lambda t,f=field:env.value(t,f),capacity)
        # This residual checks numerical boundary closure; it is not independent physics validation.
        surface=values[1:,-1]
        residual=law.integral(surface,cells[1:,-1])/mesh.half_width-transfer*(surface-ambient[1:])
        points=np.r_[0.,mesh.centres,mesh.radius]
        extended=np.column_stack((values[:,0],cells,values[:,-1]))
        linear=np.array([np.interp(cache["radius_m"],points,row) for row in extended])
        reconstruction_delta=float(np.max(np.abs(values[1:]-linear[1:])))
        assert reconstruction_delta<target
        radial_differences=np.diff(extended,axis=1)
        sign_violation=float(max(0.,np.max(-radial_differences if field=="temperature" else radial_differences)))
        assert sign_violation<1e-8
        if field=="temperature":
            assert np.max(extended-ambient[:,None])<1e-8 and np.min(extended)>=28.-1e-8
        else:
            assert np.max(extended)<=2.55+1e-8 and np.min(extended-ambient[:,None])>=-1e-8
        diagnostics["fields"][field]={
            "minimum_sampled":float(np.min(values)),"maximum_sampled":float(np.max(values)),
            "radial_monotonicity_violation_max":sign_violation,
            "surface_half_cell_flux_residual_max":float(np.max(np.abs(residual))),
            "pchip_vs_piecewise_linear_export_max":reconstruction_delta,
            "independent_balance_absolute":independent,"balance_normalization":normalizer,
            "independent_balance_relative":independent/normalizer,
            "centre_1800":float(values[-1,0]),"surface_1800":float(values[-1,-1]),
            "weighted_mean_1800":float(cache[field+"_mean"][-1]),
            "ambient_1800":float(ambient[-1])}
    dense_times=np.array([1,10,100,300,600,900,1200,1500,1800])
    dense_radii=np.linspace(0,.02,1001)
    heat_model=ScalarFVM(mesh,Diffusivity(.36),25.,lambda t:env.value(t,"temperature"),820.*2600)
    dense_heat=reconstruct(heat_model,dense_times,cache["temperature_cells"][dense_times],dense_radii)
    dense_bessel=cylinder_response(dense_times,dense_radii,env.times,env.data[:,1],28.,.36/(820*2600),.02,25*.02/.36,512)
    diagnostics["heat_bessel_dense_profiles_max_abs_C"]=float(np.max(np.abs(dense_heat-dense_bessel)))
    assert diagnostics["heat_bessel_dense_profiles_max_abs_C"]<target
    diagnostics["dense_profile_sampling"]={"times_s":dense_times.tolist(),"radial_points":1001,"is_continuous_error_bound":False}
    diagnostics["initial_water_removed_fraction_1800"]=float(1-cache["moisture_mean"][-1]/2.55)
    diagnostics["D_centre_1800_m2_s"]=float(7e-9*np.exp(-.89/cache["moisture"][-1,0]))
    diagnostics["D_surface_1800_m2_s"]=float(7e-9*np.exp(-.89/cache["moisture"][-1,-1]))
    save_json(ROOT/"reports/q1_diagnostics.json",diagnostics)
    output=ROOT/"results/final"; output.mkdir(parents=True,exist_ok=True)
    template=load_workbook(Path(cfg["source_directory"])/"result1.xlsx",read_only=True)
    template_names=template.sheetnames
    first_header=template.worksheets[0].cell(1,1).value
    template.close()
    assert template_names==list(FIELDS.values())
    workbook=Workbook();workbook.remove(workbook.active)
    workbook.properties.creator="MathModel"
    for field,chinese in FIELDS.items():
        sheet=workbook.create_sheet(chinese)
        sheet.append([first_header]+[round(i*.1,1) for i in range(21)])
        rounded=np.round(cache[field][1:],4)
        for t,row in enumerate(rounded,1):
            sheet.append([t]+row.tolist())
        sheet.freeze_panes="B2";sheet.auto_filter.ref="A1:V1801"
        sheet.column_dimensions["A"].width=29
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="254B68")
            cell.alignment=Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2,min_col=2):
            for cell in row:
                cell.number_format="0.0000"
    xlsx_path=output/"result1.xlsx"
    workbook.save(xlsx_path); workbook.close()
    # Independent I/O path: reread all cells, not a handful of sample rows.
    readback=load_workbook(xlsx_path,read_only=True,data_only=False)
    assert readback.sheetnames==template_names
    total_checked=0
    for field,chinese in FIELDS.items():
        sheet=readback[chinese]
        assert (sheet.max_row,sheet.max_column)==(1801,22)
        rows=sheet.iter_rows();header=next(rows)
        assert header[0].value==first_header
        assert [c.value for c in header[1:]]==[round(i*.1,1) for i in range(21)]
        for t,row in enumerate(rows,1):
            assert row[0].value==t
            for j,cell in enumerate(row[1:]):
                assert cell.data_type=="n" and cell.number_format=="0.0000"
                assert cell.value==float(np.round(cache[field][t,j],4)),(field,t,j,cell.value)
                total_checked+=1
    readback.close()
    export_tables(cache)
    for field in FIELDS:
        with (ROOT/f"tables/q1_{field}.csv").open(encoding="utf-8-sig",newline="") as f:
            rows=list(csv.reader(f))[1:]
        for t,row in zip(PAPER_TIMES,rows):
            assert row[0]==str(t)
            assert np.array_equal(np.array(row[1:],float),np.round(cache[field][t,PAPER_COLUMNS],4))
    np.savez_compressed(output/"q1_full_precision.npz",**{k:v for k,v in cache.items() if not k.endswith("_cells")})
    resolved={"status":"Q1_local_milestone_accepted_full_project_validation_pending",
              "model_config":"config/model_config.json","cache":"results/cache/q1/q1_final.npz",
              "parameters":{key:meta[key] for key in ("n","grading","rtol","atol","max_step_s","method","internal_flux","perturbation")},
              "time_s":[0,1800],"property_set":"q1","initialization":"fresh",
              "format_digits":4,"all_export_last_digits_proven_stable":False,
              "model_config_sha256":digest(ROOT/"config/model_config.json")}
    save_json(ROOT/"config/q1_final.json",resolved)
    verify_sources()
    paths=[xlsx_path,output/"q1_full_precision.npz",ROOT/"config/q1_final.json",
           ROOT/"tables/q1_tables.md",ROOT/"tables/q1_tables.tex",
           ROOT/"tables/q1_temperature.csv",ROOT/"tables/q1_moisture.csv",
           ROOT/"reports/q1_diagnostics.json",ROOT/"reports/kernel_checks.json",
           ROOT/"reports/q1_mesh_convergence.json",ROOT/"reports/q1_time_and_sensitivity.json",
           ROOT/"results/cache/q1/q1_final.npz",ROOT/"results/cache/q1/q1_final.json",
           ROOT/"results/cache/q1/bessel_reference.npz"]
    audit={"status":"passed_Q1_local_acceptance_and_output_audit", "numerical_cells_read_back":total_checked,
           "sheet_count":2,"data_rows_per_sheet":1800,"radius_samples_per_sheet":21,
           "original_sources_unchanged":len(original["files"]),
           "remaining_scope":["independent nonlinear weak-form spatial solver in stage5", "Q23 coupled solver", "Q4 moving material solver and tests", "full model/geometry validation"],
           "source_fingerprint":source_fingerprint(),
           "script_sha256":{str(p.relative_to(ROOT)):digest(p) for p in sorted((ROOT/"validation").glob("*.py"))+[ROOT/"scripts/question1.py",Path(__file__)]},
           "artifacts_sha256":{str(p.relative_to(ROOT)):digest(p) for p in paths}}
    assert total_checked==75600
    save_json(ROOT/"reports/q1_export_audit.json",audit)
    print(json.dumps({"audit":{k:audit[k] for k in ("status","numerical_cells_read_back","original_sources_unchanged")},"diagnostics":diagnostics},ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
