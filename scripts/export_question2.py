#!/usr/bin/env python3
"""Gate and export Q2 first-three-hour milestone; never label it full process."""
import csv
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from scripts.question2 import read_case, fingerprint
from src.boundary import Environment
from src.coupled.model import CoupledFVM
from src.coupled.integrate import reconstruct
from src.data_io import ROOT, digest, load_config, save_json, verify_sources
from src.fvm import RadialMesh


FIELDS = {"temperature": "温度", "moisture": "水分浓度"}
PAPER_TIMES = np.arange(1800, 10801, 1800)
PAPER_RADII = [0, 5, 10, 15, 20]


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def export_tables(cache):
    directory = ROOT/"tables"
    header = ["时间/h", "0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm"]
    markdown = ["# 问题二：论文表3—4", "", "附录三从原始初态重新计算；仅列前三小时。温度单位℃，含水率为kg/kg（干基）。", ""]
    table_rows = {}
    for index, (field, chinese) in enumerate(FIELDS.items(), 3):
        values = cache[field][np.ix_(PAPER_TIMES, PAPER_RADII)]
        rows = [[f"{t/3600:.1f}"]+[f"{x:.4f}" for x in row] for t, row in zip(PAPER_TIMES, values)]
        table_rows[field] = rows
        with (directory/f"q2_{field}.csv").open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(header)
            writer.writerows(rows)
        markdown += [f"## 表{index}：{chinese}", "", "| "+" | ".join(header)+" |", "|"+"---|"*6]
        markdown += ["| "+" | ".join(row)+" |" for row in rows]+[""]
    latex = [
        "% Generated from q2_final unrounded trajectory; requires booktabs and subcaption.",
        r"\begin{table}[htbp]",
        r"  \centering",
        r"  \caption{三小时内药材温湿分布}",
        r"  \label{tab:q2-results}",
    ]
    panels = [
        ("temperature", r"温度（$^\circ$C）", "q2_temperature"),
        ("moisture", r"干基含水率（kg/kg）", "q2_moisture"),
    ]
    for panel_index, (field, caption, label) in enumerate(panels):
        latex += [
            r"  \begin{subtable}[t]{0.485\textwidth}",
            r"    \centering",
            r"    \caption{"+caption+"}",
            r"    \label{tab:"+label+"}",
            r"    \fontsize{8.2pt}{10.5pt}\selectfont",
            r"    \setlength{\tabcolsep}{2.4pt}",
            r"    \begin{tabular}{@{}rrrrrr@{}}",
            r"      \toprule",
            r"      时间/h & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm \\",
            r"      \midrule",
        ]
        latex += ["      "+" & ".join(row)+r" \\" for row in table_rows[field]]
        latex += [r"      \bottomrule", r"    \end{tabular}", r"  \end{subtable}"]
        if panel_index == 0:
            latex[-1] += r"\hfill"
    latex += [r"\end{table}", ""]
    (directory/"q2_tables.md").write_text("\n".join(markdown), encoding="utf-8")
    (directory/"q2_tables.tex").write_text("\n".join(latex), encoding="utf-8")


def mechanism_diagnostics(cache, meta, model, sensitivity):
    times = cache["time_s"]
    T, C = cache["temperature"], cache["moisture"]
    ta, ce = model.ambient(times)
    b, k, _, _ = model.thermal(C)
    D = model.diffusion(C, T)[0]
    refD = float(model.diffusion(np.asarray(model.C0), np.asarray(model.T0))[0])
    points = np.r_[0., model.mesh.centres, model.mesh.radius]
    fields, _ = reconstruct(model, cache["profile_time_s"], cache["profile_state"].T, points)
    profile_air = model.ambient(cache["profile_time_s"])[0]
    air_distance = np.max(np.abs(fields[0]-profile_air[:, None]), axis=1)
    lastT, lastC = cache["final_state"][::2], cache["final_state"][1::2]
    lastH = float(np.sum(model.mesh.weights*model.thermal(lastC)[0]*(lastT-model.T0)))
    heat_out_rate = 2*model.h/model.mesh.radius*(T[:, -1]-ta)
    wrong_heat_balance = lastH+np.trapezoid(heat_out_rate, times)
    wrong_heat_balance_coarse = lastH+np.trapezoid(heat_out_rate[::2], times[::2])
    diagnostic = {
        "source_fingerprint": fingerprint(), "model_conditional_not_experimental": True,
        "temperature_3h": {"centre_C": float(T[-1, 0]), "surface_C": float(T[-1, -1]),
            "mean_C": float(cache["temperature_mean"][-1]), "air_C": float(ta[-1]),
            "air_minus_centre_C": float(ta[-1]-T[-1, 0]),
            "surface_minus_centre_C": float(T[-1, -1]-T[-1, 0]),
            "full_reconstructed_max_distance_to_air_C": float(air_distance[-1]),
            "passes_0p1C_air_distance_screen": bool(air_distance[-1] < .1)},
        "moisture_3h": {"centre": float(C[-1, 0]), "surface": float(C[-1, -1]),
            "mean": float(cache["moisture_mean"][-1]), "equivalent_air": float(ce[-1]),
            "fraction_initial_water_removed": float(1-cache["moisture_mean"][-1]/model.C0),
            "mean_rate_per_s": float(-2*model.hm/model.mesh.radius*(C[-1, -1]-ce[-1])),
            "global_max": float(cache["moisture_global_max"][-1]),
            "global_max_radius_m": float(cache["moisture_argmax_radius_m"][-1]),
            "q3_threshold_reached": bool(cache["moisture_global_max"][-1] < .15)},
        "global_radial_moisture_increase_max": float(np.max(cache["moisture_radial_increase_max"])),
        "argmax_caution": "Near-uniform initial plateaus have floating-point ties; do not interpret raw early argmax jumps physically.",
        "diffusivity_3h": {"initial_m2_s": refD,
            "centre_m2_s": float(D[-1, 0]), "surface_m2_s": float(D[-1, -1]),
            "centre_log_temperature_contribution": float(model.B*(1/(model.T0+273.15)-1/(T[-1, 0]+273.15))),
            "surface_log_temperature_contribution": float(model.B*(1/(model.T0+273.15)-1/(T[-1, -1]+273.15))),
            "centre_log_moisture_contribution": float(model.a*(1/model.C0-1/C[-1, 0])),
            "surface_log_moisture_contribution": float(model.a*(1/model.C0-1/C[-1, -1])),
            "alpha_over_D_centre": float(k[-1, 0]/b[-1, 0]/D[-1, 0]),
            "alpha_over_D_surface": float(k[-1, -1]/b[-1, -1]/D[-1, -1])},
        "incorrect_uncorrected_heat_balance_diagnostic": {
            "uncorrected_residual_J_m3": float(wrong_heat_balance),
            "trapezoid_1s_vs_2s_difference_J_m3": float(wrong_heat_balance-wrong_heat_balance_coarse),
            "is_not_the_valid_heat_conservation_residual": True},
        "profile_air_distance": {"time_s": cache["profile_time_s"].tolist(), "max_abs_C": air_distance.tolist()},
        "sensitivity": []}
    for row in sensitivity["cases"]:
        alternative, _ = read_case(row["label"])
        diagnostic["sensitivity"].append({"label": row["label"], "options": row["options"],
            "max_temperature_change_C": float(np.max(np.abs(alternative["temperature"]-T))),
            "max_moisture_change": float(np.max(np.abs(alternative["moisture"]-C))),
            "mean_moisture_3h": float(alternative["moisture_mean"][-1]),
            "initial_water_removed_fraction_3h": float(1-alternative["moisture_mean"][-1]/model.C0)})
    return diagnostic


def main():
    original = verify_sources()
    cfg = load_config()
    cache, meta = read_case("q2_final")
    convergence = read_json(ROOT/"reports/q2_convergence.json")
    sensitivity = read_json(ROOT/"reports/q2_sensitivity.json")
    independent = read_json(ROOT/"reports/q2_independent_verification.json")
    kernel = read_json(ROOT/"reports/q2_kernel_checks.json")
    for report in (convergence, sensitivity, independent, kernel):
        assert report["source_fingerprint"] == fingerprint()
    assert all(r["complete"] for r in (convergence, sensitivity, independent))
    assert kernel["status"].startswith("passed_") and independent["status"].startswith("passed_")
    assert len(sensitivity["cases"]) == 10
    assert (meta["n"], meta["rtol"], meta["max_step_s"], meta["perturbation"]) == (1280, 1e-10, 2., {})
    target = cfg["diagnostics"]["numerical_absolute_target_for_four_decimal_claim"]
    for row in convergence["mesh_series"]:
        read_case(row["label"])
        if row["n"] >= 640:
            assert all(d["max_abs"] < target for d in row["difference_from_previous"].values())
    for row in convergence["temporal"]:
        _, candidate_meta = read_case(row["label"])
        if row["label"] in ("q2_cap4", "q2_cap1", "q2_radau"):
            assert all(d["max_abs"] < target for d in row["difference_from_baseline"].values())
        if row["label"] in ("q2_cap4", "q2_cap1"):
            assert abs(candidate_meta["statistics"]["accepted_max_step_s"]-candidate_meta["max_step_s"]) < 1e-8
            assert candidate_meta["statistics"]["accepted_steps"] != meta["statistics"]["accepted_steps"]
    assert abs(meta["statistics"]["accepted_max_step_s"]-2.) < 1e-8
    for row in independent["backward_euler"]+independent["bdf2"]:
        folder = ROOT/"results/cache/q2_independent"
        im = read_json(folder/(row["label"]+".json"))
        assert im["script_sha256"] == digest(ROOT/"validation/q2_independent_fem.py")
        assert im["cache_sha256"] == digest(folder/(row["label"]+".npz"))
    assert meta["statistics"]["water_balance_relative"] < 1e-6
    assert meta["statistics"]["heat_corrected_balance_relative"] < 1e-6
    assert np.array_equal(cache["time_s"], np.arange(10801.))
    assert np.array_equal(cache["profile_state"][0, ::2], np.full(meta["n"], 28.))
    assert np.array_equal(cache["profile_state"][0, 1::2], np.full(meta["n"], 2.55))
    for field in FIELDS:
        assert cache[field].shape == (10801, 21) and np.isfinite(cache[field]).all()
    assert np.min(cache["moisture"]) > 0 and np.max(cache["moisture"]) <= 2.55+1e-10
    assert np.max(cache["moisture_radial_increase_max"]) < 1e-9
    assert cache["moisture_global_max"][-1] > .15
    model = CoupledFVM(RadialMesh(meta["n"], .02, meta["grading"]), cfg, Environment(cfg))
    diagnostic = mechanism_diagnostics(cache, meta, model, sensitivity)
    save_json(ROOT/"reports/q2_mechanisms.json", diagnostic)
    export_tables(cache)
    output = ROOT/"results/alternatives"
    output.mkdir(parents=True, exist_ok=True)
    template = load_workbook(Path(cfg["source_directory"])/"result2.xlsx", read_only=True)
    assert template.sheetnames == list(FIELDS.values())
    header = template.worksheets[0].cell(1, 1).value
    template.close()
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "MathModel"
    workbook.properties.description = "Q2 first 10800 seconds only. Full-process result2.xlsx awaits Q3 event."
    for field, chinese in FIELDS.items():
        sheet = workbook.create_sheet(chinese)
        sheet.freeze_panes = "B2"
        sheet.append([header]+[round(j*.1, 1) for j in range(21)])
        for t in range(1, 10801):
            cells = [t]
            for value in np.round(cache[field][t], 4):
                cell = WriteOnlyCell(sheet, value=float(value))
                cell.number_format = "0.0000"
                cells.append(cell)
            sheet.append(cells)
    path = output/"result2_first3h.xlsx"
    workbook.save(path)
    workbook.close()
    readback = load_workbook(path, read_only=True, data_only=False)
    assert readback.sheetnames == list(FIELDS.values())
    count = 0
    for field, chinese in FIELDS.items():
        rows = readback[chinese].iter_rows()
        first = next(rows)
        assert first[0].value == header
        assert [c.value for c in first[1:]] == [round(j*.1, 1) for j in range(21)]
        for t, row in enumerate(rows, 1):
            assert row[0].value == t and len(row) == 22
            for j, cell in enumerate(row[1:]):
                assert cell.data_type == "n" and cell.number_format == "0.0000"
                assert cell.value == float(np.round(cache[field][t, j], 4))
                count += 1
        assert t == 10800
    readback.close()
    assert count == 453600
    for field in FIELDS:
        with (ROOT/f"tables/q2_{field}.csv").open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))[1:]
        for t, row in zip(PAPER_TIMES, rows):
            assert float(row[0])*3600 == t
            assert np.array_equal(np.asarray(row[1:], float), np.round(cache[field][t, PAPER_RADII], 4))
    full_precision = output/"q2_first3h_full_precision.npz"
    np.savez_compressed(full_precision, **cache)
    final_config = {"status": "Q2_first3h_local_and_independent_sample_checks_passed_full_process_pending_Q3",
        "property_set": "q23", "initialization": "fresh", "physical_time_origin_s": 0.,
        "time_s": [0., 10800.], "all_export_last_digits_proven_stable": False,
        "parameters": {k: meta[k] for k in ("n", "grading", "rtol", "atol_temperature", "atol_moisture", "max_step_s", "method", "perturbation", "boundary_half_cell")},
        "source_fingerprint": fingerprint(), "continuation_checkpoint": "results/cache/q23/q2_final.npz:final_state",
        "full_result2_xlsx_generated": False}
    save_json(ROOT/"config/q2_first3h.json", final_config)
    verify_sources()
    paths = [path, full_precision, ROOT/"config/q2_first3h.json", ROOT/"reports/q2_mechanisms.json",
             ROOT/"reports/q2_convergence.json", ROOT/"reports/q2_sensitivity.json",
             ROOT/"reports/q2_kernel_checks.json", ROOT/"reports/q2_independent_verification.json",
             ROOT/"tables/q2_tables.md", ROOT/"tables/q2_tables.tex",
             ROOT/"tables/q2_temperature.csv", ROOT/"tables/q2_moisture.csv"]
    audit = {"status": "passed_Q2_first3h_export_and_local_numerical_acceptance",
        "source_fingerprint": fingerprint(), "numerical_cells_read_back": count,
        "original_sources_unchanged": len(original["files"]), "full_process_workbook_pending_Q3": True,
        "independent_comparison_is_sampled_not_continuous_error_bound": True,
        "script_sha256": {str(p.relative_to(ROOT)): digest(p) for p in
            sorted((ROOT/"validation").glob("q2*.py"))+[Path(__file__), ROOT/"scripts/question2.py"]},
        "artifacts_sha256": {str(p.relative_to(ROOT)): digest(p) for p in paths}}
    save_json(ROOT/"reports/q2_export_audit.json", audit)
    print(json.dumps({"status": audit["status"], "checked_values": count,
                      "temperature_3h": diagnostic["temperature_3h"], "moisture_3h": diagnostic["moisture_3h"]}, indent=2))


if __name__ == "__main__":
    main()
