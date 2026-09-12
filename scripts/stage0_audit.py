#!/usr/bin/env python3
"""Reproduce stage-0 input checks and algebraic diagnostics; never solve the PDE.

Generated reports contain data statistics, local coefficient scales and identities.
Neither radius-implied moisture nor diffusion time scales are drying predictions.
"""
import csv
import copy
import hashlib
import importlib.metadata
import json
import math
import platform
import sys
from pathlib import Path

import numpy as np
import sympy as sp
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
CONFIG_PATH = ROOT / "config/model_config.json"


def save_json(name, value):
    (REPORTS / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def file_digest(path):
    with path.open("rb") as stream:
        digest = hashlib.sha256()
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_audit(path):
    workbook = load_workbook(path, data_only=False, read_only=True)
    result = {}
    for sheet in workbook:
        rows = list(sheet.values)
        result[sheet.title] = {
            "rows_including_header": len(rows),
            "columns": sheet.max_column,
            "header": list(rows[0]),
            "nonempty_cells": sum(v is not None for row in rows for v in row),
            "formula_cells": sum(isinstance(v, str) and v.startswith("=") for row in rows for v in row),
        }
    workbook.close()
    return result


def input_table(path, expected_columns, expected_count, expected_dt):
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = list(workbook.active.values)
    workbook.close()
    assert list(rows[0]) == expected_columns
    assert len(rows) - 1 == expected_count
    assert all(v is not None for row in rows[1:] for v in row)
    values = np.asarray(rows[1:], dtype=float)
    assert np.isfinite(values).all()
    assert np.all(np.diff(values[:, 0]) == expected_dt)
    assert len(np.unique(values[:, 0])) == expected_count
    return values


def properties(config, group, C, T_C):
    if C <= 0 or T_C + 273.15 <= 0:
        raise ValueError("Stage0 coefficient diagnostics require C>0 and T_K>0")
    p = config["property_sets"][group]
    rho = p["rho_kg_m3"]["offset"] + p["rho_kg_m3"]["slope_C"] * C
    cp = p["cp_J_kg_K"]["offset"] + p["cp_J_kg_K"]["fraction_C_over_1_plus_C"] * C / (1 + C)
    k = p["k_W_m_K"]["offset"] + p["k_W_m_K"]["fraction_C_over_1_plus_C"] * C / (1 + C)
    d = p["D_m2_s"]
    D = d["prefactor"] * math.exp(-d["moisture_exponent"] / C - d["thermal_exponent_K"] / (T_C + 273.15))
    return rho, cp, k, D


def check_properties(config):
    """Independent literal appendix checks catch configuration transcription errors."""
    tests = []
    for C, T in [(2.55, 28), (0.15, 50), (0.05, 50), (1.0, 40)]:
        expected = {
            "q1": (820, 2600, .36, 7e-9 * math.exp(-.89/C)),
            "q23": (650+128*C, 1450+2736*C/(1+C), .21+.38*C/(1+C), 2.4e-3*math.exp(-.45/C)*math.exp(-3850/(T+273.15))),
            "q4": (760+90*C, 1850+2150*C/(1+C), .12+.20*C/(1+C), 4.2e-4*math.exp(-.30/C)*math.exp(-3850/(T+273.15))),
        }
        for group, reference in expected.items():
            computed = properties(config, group, C, T)
            assert np.allclose(computed, reference, rtol=1e-13, atol=0)
            tests.append({"group": group, "C": C, "T_C": T, "passed": True})
    return tests


def scales(config):
    R0 = config["geometry"]["reference_radius_m"]
    h = config["boundary"]["heat_transfer_W_m2_K"]
    hm = config["boundary"]["normalized_mass_transfer_m_s"]
    rows = []
    for group in ["q1", "q23", "q4"]:
        for C, T in [(2.55, 28), (2.55, 50), (1.0, 50), (.15, 50), (.05, 50)]:
            rho, cp, k, D = properties(config, group, C, T)
            alpha = k/(rho*cp)
            rows.append({
                "properties": group, "C": C, "T_C": T, "diagnostic_radius_m": R0,
                "rho_eff_kg_m3": rho, "cp_J_kg_K": cp, "k_W_m_K": k,
                "alpha_m2_s": alpha, "D_m2_s": D,
                "tau_T_min": R0**2/alpha/60, "tau_C_h": R0**2/D/3600,
                "tau_heat_external_min": rho*cp*R0/(2*h)/60,
                "tau_mass_external_h": R0/(2*hm)/3600,
                "Bi_T_R": h*R0/k, "Bi_C_R": hm*R0/D,
                "alpha_over_D": alpha/D,
                "epsilon_D_over_alpha": D/alpha,
                "d_lnD_d_C": config["property_sets"][group]["D_m2_s"]["moisture_exponent"]/C**2,
                "d_lnD_d_TK": config["property_sets"][group]["D_m2_s"]["thermal_exponent_K"]/(T+273.15)**2,
            })
    with (REPORTS/"scales_snapshot.csv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return rows


def symbolic_checks():
    """Verify algebraic identities and analytic fields without a PDE solver."""
    r, x, t = sp.symbols("r xi t", real=True)
    R = sp.Function("R")(t)
    # An explicit nonseparable smooth field avoids unevaluated Subs(Derivative)
    # canonicalization differences; this checks the transform, not a PDE solution.
    C = sp.exp(r*t) + r**2*sp.sin(t)
    u = C.subs(r, R*x)
    material = (sp.diff(C,t) + sp.diff(R,t)/R*r*sp.diff(C,r)).subs(r,R*x)
    checks = {"material_chain_rule_analytic_profile": sp.simplify(sp.diff(u,t)-material) == 0}
    s0, R0 = sp.symbols("s0 R0", positive=True)
    s = s0*(R0/R)**2
    velocity = sp.diff(R,t)/R*r
    checks["dry_solid_continuity"] = sp.simplify(sp.diff(s,t) + sp.diff(r*s*velocity,r)/r) == 0
    v, w, ut, ux = sp.symbols("v w ut ux")
    ct = ut-w*ux/R
    checks["general_relative_velocity"] = sp.simplify(ct+v*ux/R-(ut+(v-w)*ux/R)) == 0

    # Nonconstant diffusivity and profile exercise all derivative terms.
    f = 2 + r**2/R**2 + sp.exp(-t)*r**4/R**4
    D = 1 + f
    physical_op = sp.diff(r*D*sp.diff(f,r),r)/r
    U = f.subs(r,R*x)
    ref_op = sp.diff(x*(1+U)*sp.diff(U,x),x)/(R**2*x)
    checks["radial_operator_mapping"] = sp.simplify(physical_op.subs(r,R*x)-ref_op) == 0
    a0, a2, b0, b2 = sp.symbols("a0 a2 b0 b2")
    center_op = sp.diff(r*(a0+a2*r**2)*sp.diff(b0+b2*r**2,r),r)/r
    checks["axis_limit_2a_frr"] = sp.simplify(sp.limit(center_op,r,0)-4*a0*b2) == 0
    reference_mass = s*R**2/2
    checks["dry_mass_constant_under_shrinkage"] = sp.simplify(sp.diff(reference_mass,t)) == 0
    U_static = 1+x**2
    water_mass = s*R**2*sp.integrate(x*U_static,(x,0,1))
    checks["nonuniform_material_water_mass_constant_no_flux"] = sp.simplify(sp.diff(water_mass,t)) == 0
    c = sp.symbols("c", nonnegative=True)
    density_derivative = sp.diff((760+90*c)/(1+c),c)
    checks["strict_density_ratio_decreasing"] = sp.simplify(density_derivative+670/(1+c)**2) == 0
    # Celsius state and Kelvin exponent derivative have identical differential increments.
    temp, exponent = sp.symbols("temp exponent", positive=True)
    checks["arrhenius_log_derivative"] = sp.simplify(sp.diff(-exponent/(temp+sp.Rational(27315,100)),temp)-exponent/(temp+sp.Rational(27315,100))**2) == 0
    assert all(checks.values()), checks
    return checks


def validate_contract(config):
    p = config["problems"]
    assert p["q1"]["properties"] == "q1" and p["q1"]["initialization"] == "fresh"
    assert p["q2"]["properties"] == "q23" and p["q2"]["initialization"] == "fresh"
    assert p["q3"]["properties"] == p["q2"]["properties"] and p["q3"]["time_origin_s"] == 0
    assert p["q4"]["properties"] == "q4" and p["q4"]["initialization"] == "fresh"
    assert config["closure"]["empirical_density_is_not_transport_dry_density"] is True
    assert config["geometry"]["q4_mesh_motion"] == "follows_material"
    assert config["events"]["direction"] == -1


def main():
    REPORTS.mkdir(parents=True, exist_ok=True)
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    validate_contract(config)
    rejected = []
    for problem, key, bad_value in [
        ("q2", "initialization", "q1_end"),
        ("q2", "properties", "q1"),
        ("q3", "time_origin_s", 10800),
        ("q4", "initialization", "q3_end"),
    ]:
        bad_config = copy.deepcopy(config)
        bad_config["problems"][problem][key] = bad_value
        try:
            validate_contract(bad_config)
        except AssertionError:
            rejected.append({"problem":problem,"field":key,"invalid_value":bad_value,"rejected":True})
        else:
            raise AssertionError("invalid timeline/config accepted")
    base = Path(config["source_directory"])
    paths = [base/name for name in ["A题.pdf","附件1.xlsx","附件2.xlsx","result1.xlsx","result2.xlsx","result3.xlsx","result4.xlsx"]]
    paths.append(Path(config["research_report"]))
    input_hashes = {str(p): file_digest(p) for p in paths}
    records = [{"path":str(p),"bytes":p.stat().st_size,"sha256":input_hashes[str(p)]} for p in paths]
    A = input_table(base/"附件1.xlsx", ["时间","温度","水分浓度"],241,60)
    B = input_table(base/"附件2.xlsx", ["时间","半径"],145,1800)
    assert A[0,0] == 0 and A[-1,0] == 14400
    assert B[0,0] == 0 and B[-1,0] == 259200
    assert B[0,1] == 2 and (B[:,1] > 0).all() and (np.diff(B[:,1]) <= 0).all()
    assert A[0,1] == config["initial"]["temperature_C"]
    for n in range(1,5):
        wb = load_workbook(base/f"result{n}.xlsx",read_only=True,data_only=False)
        expected = ["温度","水分浓度"] if n<3 else ["Sheet1"]
        assert wb.sheetnames == expected
        for sheet in wb:
            assert all(cell.value is None for row in sheet.iter_rows(min_row=2,min_col=2) for cell in row)
        wb.close()
    windows = {}
    for name,lo,hi in [("main_3_4h",10800,14400),("alternative_2_5_4h",9000,14400)]:
        sample=A[(A[:,0]>=lo)&(A[:,0]<=hi)]
        windows[name] = {"window_s":[lo,hi],"count":len(sample),"T_mean_C":float(sample[:,1].mean()),"Y_mean":float(sample[:,2].mean()),"T_std_C":float(sample[:,1].std(ddof=1)),"Y_std":float(sample[:,2].std(ddof=1))}
    windows["last_observation"]={"T_C":float(A[-1,1]),"Y":float(A[-1,2])}

    def first_fraction(col, initial, plateau, fraction):
        threshold=initial+fraction*(plateau-initial)
        selected=A[A[:,col]>=threshold]
        return {"threshold":float(threshold),"first_observed_time_s":int(selected[0,0]),"meaning":"first crossing in sampled noisy data, not fitted time constant"}
    environment_rise={}
    for name,col in [("temperature",1),("air_moisture",2)]:
        plateau=windows["main_3_4h"]["T_mean_C" if col==1 else "Y_mean"]
        environment_rise[name]={str(frac):first_fraction(col,A[0,col],plateau,frac) for frac in [.5,.9,.95]}

    stage0_scales=scales(config)
    rho0=(760+90*2.55)
    s0=rho0/3.55
    rmin_cm=2*math.sqrt(s0/760)
    required_last=s0*(2/B[-1,1])**2
    first_bad=B[B[:,1]<rmin_cm][0]
    density={
        "assumptions_being_tested":["rho(C) is exact wet bulk mass density","constant length","dry solid conserved","C>=0"],
        "initial_wet_density_kg_m3":rho0,"initial_dry_density_kg_m3":s0,
        "maximum_dry_density_if_rho_exact_kg_m3":760,
        "necessary_minimum_radius_cm":rmin_cm,
        "observed_final_radius_cm":float(B[-1,1]),
        "required_final_dry_density_kg_m3":float(required_last),
        "first_sample_violating_necessary_bound_s":int(first_bad[0]),
        "first_sample_radius_cm":float(first_bad[1]),
        "implied_uniform_C_final_if_all_assumptions_forced":float((760-required_last)/(required_last-90)),
        "maximum_final_dry_mass_over_initial_if_rho_exact":float(760*(B[-1,1]/2)**2/s0),
        "interpretation":"incompatibility of simultaneous strict assumptions, not evidence that input data are wrong",
        "adopted_closure":config["closure"]
    }
    assert required_last>760
    Rmid=(B[:-1,1]+B[1:,1])/200
    Rdot=np.diff(B[:,1])/100/np.diff(B[:,0])
    shrink=[]
    for i in range(len(Rmid)):
        shrink.append({"t_start_s":int(B[i,0]),"t_end_s":int(B[i+1,0]),"R_mid_m":float(Rmid[i]),"R_dot_m_s":float(Rdot[i]),"tau_sh_h":None if Rdot[i]==0 else float(Rmid[i]/abs(Rdot[i])/3600)})
    save_json("shrinkage_scales.json",shrink)
    symbolic=symbolic_checks()
    property_tests=check_properties(config)
    audit={
        "scope":"stage0_no_pde_solution",
        "pdf_pages":len(PdfReader(base/"A题.pdf").pages),
        "workbooks":{p.name:workbook_audit(p) for p in paths if p.suffix==".xlsx"},
        "attachment1":{"records":241,"time_interval_s":60,"time_end_s":14400,"temperature_range_C":[float(A[:,1].min()),float(A[:,1].max())],"air_moisture_range":[float(A[:,2].min()),float(A[:,2].max())],"missing":0,"duplicate_times":0},
        "attachment2":{"records":145,"time_interval_s":1800,"time_end_s":259200,"radius_range_cm":[float(B[:,1].min()),float(B[:,1].max())],"missing":0,"duplicate_times":0,"monotone_nonincreasing":True},
        "environment_extension":windows,"environment_rise_diagnostics":environment_rise,
        "initial_shrinkage_interval":shrink[0],
        "density_consistency":density
    }
    save_json("data_audit.json",audit)
    save_json("density_consistency.json",density)
    save_json("resolved_boundary.json",windows)
    save_json("stage0_checks.json",{"status":"passed_algebra_and_input_checks_only","property_checks":property_tests,"symbolic_checks":symbolic,"config_contract":"passed","invalid_timeline_configs":rejected,"pde_solver_tests":"not_run_solver_not_implemented","physical_validation":"not_available_no_internal_measurements"})
    dependencies={name:importlib.metadata.version(name) for name in ["numpy","scipy","openpyxl","pypdf","sympy","et-xmlfile","mpmath","typing-extensions"]}
    save_json("environment.json",{"python":sys.version,"executable":sys.executable,"platform":platform.platform(),"base_prefix":sys.base_prefix,"prefix":sys.prefix,"isolated_virtualenv":sys.prefix!=sys.base_prefix,"dependencies":dependencies})
    for p in paths:
        assert file_digest(p)==input_hashes[str(p)],"input changed during stage0: "+str(p)
    save_json("input_manifest.json",{"files":records,"config_sha256":file_digest(CONFIG_PATH),"audit_script_sha256":file_digest(Path(__file__)),"sources_unchanged_after_run":True})
    print(json.dumps({"data_checks":"passed","property_cases":len(property_tests),"symbolic_checks":len(symbolic),"pde_solved":False,"environment_main":windows["main_3_4h"],"Q4_density_Rmin_cm":rmin_cm,"Q4_density_first_incompatible_sample_h":float(first_bad[0]/3600),"initial_scale_rows":[r for r in stage0_scales if r["C"]==2.55 and r["T_C"]==28]},ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
