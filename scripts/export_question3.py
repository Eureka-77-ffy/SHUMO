#!/usr/bin/env python3
"""Gate Q3 and export complete Q2/Q3 workbooks from one unrounded trajectory."""
import csv
import json
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from scripts.question2 import read_case as read_q2
from scripts.question3 import read_case, fingerprint
from src.coupled.model import CoupledFVM
from src.coupled.integrate import reconstruct
from src.data_io import ROOT, digest, load_config, save_json, verify_sources
from src.events.integration import ScenarioEnvironment
from src.fvm import RadialMesh
from validation.q3_independent_fem import fingerprint as independent_fingerprint


def read_json(path):
    return json.loads(Path(path).read_text())


def write_and_read(path, names, fields, times, header):
    wb=Workbook(write_only=True)
    wb.properties.creator="MathModel"
    wb.properties.description="Unrounded Q23 trajectory; terminal row is the critical crossing, not strict Cmax < 0.15."
    for name,values in zip(names,fields):
        sheet=wb.create_sheet(name)
        sheet.freeze_panes="B2"
        sheet.append([header]+[round(j*.1,1) for j in range(21)])
        for i,(t,row) in enumerate(zip(times,values),1):
            tc=WriteOnlyCell(sheet,value=float(t))
            tc.number_format="0" if t==int(t) else "0.0000"
            cells=[tc]
            for value in np.round(row,4):
                cell=WriteOnlyCell(sheet,value=float(value))
                cell.number_format="0.0000"
                cells.append(cell)
            sheet.append(cells)
            if i%50000==0:
                print(f"{path.name} {name}: wrote {i} rows",flush=True)
    wb.save(path); wb.close()
    print(f"{path.name}: saved, reading back every numeric cell",flush=True)
    wb=load_workbook(path,read_only=True,data_only=False)
    assert wb.sheetnames==names
    count=0
    for name,values in zip(names,fields):
        rows=wb[name].iter_rows()
        first=next(rows)
        assert first[0].value==header
        assert [c.value for c in first[1:]]==[round(j*.1,1) for j in range(21)]
        for i,row in enumerate(rows):
            assert i<len(times) and len(row)==22 and abs(row[0].value-times[i])<1e-8
            expected=np.round(values[i],4)
            for j,cell in enumerate(row[1:]):
                assert cell.data_type=="n" and cell.number_format=="0.0000"
                assert cell.value==float(expected[j]),(name,i,j)
                count+=1
            if (i+1)%50000==0:
                print(f"{path.name} {name}: verified {i+1} rows",flush=True)
        assert i+1==len(times)
    wb.close()
    assert count==len(times)*21*len(names)
    return count


def mechanism_report(full,tail,meta,model,sensitivity):
    st=meta["statistics"]
    T,C=full["temperature"][-1],full["moisture"][-1]
    ta,ce=model.ambient(st["event_time_s"])
    D=model.diffusion(C,T)[0]
    b,k,_,_=model.thermal(C)
    points=np.r_[0.,model.mesh.centres,.02]
    fields,_=reconstruct(model,tail["profile_time_s"],tail["profile_state"].T,points)
    air=model.ambient(tail["profile_time_s"])[0]
    air_distance=np.max(np.abs(fields[0]-air[:,None]),axis=1)
    passed=np.flatnonzero(air_distance<.1)
    retained_screen=None
    for i in passed:
        if np.all(air_distance[i:]<.1) and tail["profile_time_s"][-1]-tail["profile_time_s"][i]>=1800.:
            retained_screen=float(tail["profile_time_s"][i]); break
    rows=[]
    for row in sensitivity["cases"]:
        rows.append({"label":row["label"],"perturbation":row["perturbation"],
            "boundary_scenario":row["boundary_scenario"],"event_h":row["event_time_s"]/3600,
            "event_difference_s":row["event_time_s"]-st["event_time_s"],
            "event_relative_difference":row["event_time_s"]/st["event_time_s"]-1})
    diagnostic={"model_conditional_not_experimental":True,"source_fingerprint":fingerprint(),
        "event_time_s":st["event_time_s"],"event_time_h":st["event_time_s"]/3600,
        "threshold":.15,"critical_centre_moisture":float(C[0]),"surface_moisture":float(C[-1]),
        "mean_moisture":float(full["moisture_mean"][-1]),
        "initial_water_removed_fraction":float(1-full["moisture_mean"][-1]/model.C0),
        "centre_temperature_C":float(T[0]),"surface_temperature_C":float(T[-1]),
        "air_temperature_C":float(ta),"equivalent_air_moisture":float(ce),
        "centre_D_m2_s":float(D[0]),"surface_D_m2_s":float(D[-1]),
        "centre_to_surface_D_ratio":float(D[0]/D[-1]),
        "alpha_over_D_centre":float(k[0]/b[0]/D[0]),"alpha_over_D_surface":float(k[-1]/b[-1]/D[-1]),
        "internal_moisture_drop_fraction":float((C[0]-C[-1])/(C[0]-ce)),
        "normalized_surface_water_flux_m_s":float(model.hm*(C[-1]-ce)),
        "mean_moisture_rate_per_s":float(-2*model.hm/.02*(C[-1]-ce)),
        "event_log_D_temperature_contributions":[float(model.B*(1/(model.T0+273.15)-1/(value+273.15))) for value in (T[0],T[-1])],
        "event_log_D_moisture_contributions":[float(model.a*(1/model.C0-1/value)) for value in (C[0],C[-1])],
        "all_tail_argmax_at_centre":bool(np.all(tail["moisture_argmax_radius_m"]==0)),
        "argmax_was_searched_not_assumed":True,
        "screened_near_isothermal_since_sampled_time_s":retained_screen,
        "isothermal_screen_does_not_replace_temperature_PDE":True,
        "profile_air_distance":{"time_s":tail["profile_time_s"].tolist(),"max_abs_C":air_distance.tolist()},
        "sensitivity":rows}
    return diagnostic


def main():
    original=verify_sources(); cfg=load_config()
    tail,meta=read_case("q3_final")
    prefix,pm=read_q2(meta["q2_parent_label"])
    reports={name:read_json(ROOT/f"reports/q3_{name}.json") for name in
             ("convergence","independent_verification","event_checks","sensitivity")}
    for name,r in reports.items():
        assert r["source_fingerprint"]==fingerprint()
        assert r["script_sha256"]==digest(ROOT/f"validation/q3_{name}.py")
        if name!="event_checks": assert r["complete"]
    assert reports["event_checks"]["status"].startswith("passed_")
    assert len(reports["sensitivity"]["cases"])==12
    convergence=reports["convergence"]
    for r in convergence["mesh_series"]+convergence["temporal"]+reports["sensitivity"]["cases"]:
        read_case(r["label"])
    assert abs(convergence["mesh_series"][-1]["event_delta_s"])<.18
    for r in convergence["temporal"]:
        if r["label"] in ("q3_cap60","q3_cap15","q3_radau"):
            assert abs(r["event_delta_s"])<.18
            assert all(d["max_abs"]<5e-5 for d in r["difference_from_baseline"]["fields"].values())
    independent=reports["independent_verification"]
    assert independent["independent_fingerprint"]==independent_fingerprint()
    for r in independent["cases"]:
        folder=ROOT/"results/cache/q3_independent"
        im=read_json(folder/(r["label"]+".json"))
        assert im["source_fingerprint"]==independent_fingerprint()
        assert im["cache_sha256"]==digest(folder/(r["label"]+".npz"))
        assert im["initial_parent_sha256"]==digest(ROOT/"results/cache/q2_independent"/(im["initial_parent_label"]+".npz"))
    fine=independent["cases"][-1]
    assert fine["elements"]==2560 and abs(fine["event_delta_from_primary_s"])<.18
    assert all(d<5e-5 for d in fine["sampled_field_max_difference"].values())
    assert (meta["n"],meta["max_step_s"],meta["output_step_s"],meta["perturbation"],meta["boundary_scenario"])==(1280,30.,1.,{},"main")
    assert np.array_equal(tail["profile_state"][0],prefix["final_state"])
    st=meta["statistics"]
    assert st["water_balance_relative"]<1e-6 and st["heat_corrected_balance_relative"]<1e-6
    assert st["global_max_temporal_increase_at_checked_points"]<1e-10
    full={"radius_m":prefix["radius_m"]}
    for key in ("time_s","temperature","moisture","temperature_mean","moisture_mean",
                "moisture_global_max","moisture_argmax_radius_m","moisture_radial_increase_max"):
        full[key]=np.concatenate((prefix[key],tail[key][1:]),axis=0)
    event=st["event_time_s"]
    assert np.array_equal(full["time_s"][:-1],np.arange(np.floor(event)+1))
    assert full["time_s"][-1]==event
    assert np.all(full["moisture_global_max"][:-1]>.15)
    assert abs(full["moisture_global_max"][-1]-.15)<1e-10
    assert np.max(full["moisture_radial_increase_max"])<1e-9
    model=CoupledFVM(RadialMesh(1280,.02,2.),cfg,ScenarioEnvironment(cfg))
    diagnostic=mechanism_report(full,tail,meta,model,reports["sensitivity"])
    diagnostic["whole_process_balance_relative_upper_bounds_from_prefix_plus_tail"]={
        "water":pm["statistics"]["water_balance_relative"]+st["water_balance_relative"],
        "corrected_heat":pm["statistics"]["heat_corrected_balance_relative"]+st["heat_corrected_balance_relative"]}
    save_json(ROOT/"reports/q3_mechanisms.json",diagnostic)
    output=ROOT/"results/final"; output.mkdir(parents=True,exist_ok=True)
    np.savez_compressed(output/"q23_full_precision.npz",**full)
    paper_t=np.r_[np.arange(21600.,event,21600.),event]
    paper_i=np.r_[paper_t[:-1].astype(int),len(full["time_s"])-1]
    columns=[0,5,10,15,20]
    rows=[[f"{t/3600:.4f}"]+[f"{v:.4f}" for v in row] for t,row in
          zip(paper_t,full["moisture"][np.ix_(paper_i,columns)])]
    header=["时间/h","0 cm","0.5 cm","1 cm","1.5 cm","2 cm"]
    with (ROOT/"tables/q3_moisture.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f); writer.writerow(header); writer.writerows(rows)
    md=["# 问题三：论文表5","","由Q2同一全精度轨迹生成；末行为临界点 Cmax=0.15，并非严格小于阈值。","",
        "| "+" | ".join(header)+" |","|"+"---|"*6]+["| "+" | ".join(r)+" |" for r in rows]
    (ROOT/"tables/q3_tables.md").write_text("\n".join(md)+"\n",encoding="utf-8")
    tex=[r"\begin{table}[htbp]",r"\centering",r"\caption{全过程径向干基含水率}",
         r"\label{tab:q3_moisture}",r"\begin{tabular}{rrrrrr}",r"\toprule",
         r"时间/h & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm \\",r"\midrule"]
    tex += [" & ".join(row)+r" \\" for row in rows]
    tex += [r"\bottomrule",r"\end{tabular}",r"\end{table}",""]
    (ROOT/"tables/q3_tables.tex").write_text("\n".join(tex),encoding="utf-8")
    counts={}
    for question in (2,3):
        wb=load_workbook(Path(cfg["source_directory"])/f"result{question}.xlsx",read_only=True)
        names=wb.sheetnames; first_header=wb.worksheets[0].cell(1,1).value; wb.close()
        indices=np.arange(1,len(full["time_s"])) if question==2 else np.r_[np.arange(60,int(event)+1,60),len(full["time_s"])-1]
        fields=[full["temperature"][indices],full["moisture"][indices]] if question==2 else [full["moisture"][indices]]
        assert names==(["温度","水分浓度"] if question==2 else ["Sheet1"])
        counts[f"result{question}.xlsx"]=write_and_read(output/f"result{question}.xlsx",names,fields,full["time_s"][indices],first_header)
    verify_sources()
    # Earlier Q1/Q2 files must remain byte-for-byte unchanged.
    for oldname in ("q1_export_audit.json","q2_export_audit.json"):
        old=read_json(ROOT/"reports"/oldname)
        for relative,h in old["artifacts_sha256"].items():
            assert digest(ROOT/relative)==h,relative
    resolved={"status":"Q23_numerical_and_independent_event_milestone_accepted_model_validity_pending",
        "q2_parent_label":meta["q2_parent_label"],"q3_tail_label":"q3_final",
        "property_set":"q23","physical_time_origin_s":0.,"critical_time_s":event,
        "critical_time_h":event/3600,"first_checked_strict_integer_second":st["first_checked_integer_second"],
        "first_strict_continuous_minimum_exists":False,"all_export_last_digits_proven_stable":False,
        "source_fingerprint":fingerprint(),"parameters":{k:meta[k] for k in
           ("n","grading","rtol","atol_temperature","atol_moisture","max_step_s","boundary_scenario")},
        "prefix_max_step_s":pm["max_step_s"],"extension":"mean of observations 3-4 h, held after 4 h",
        "full_result2_xlsx_generated":True,"full_result3_xlsx_generated":True}
    save_json(ROOT/"config/q23_final.json",resolved)
    paths=[output/"result2.xlsx",output/"result3.xlsx",output/"q23_full_precision.npz",
           ROOT/"config/q23_final.json",ROOT/"reports/q3_mechanisms.json",
           ROOT/"tables/q3_tables.md",ROOT/"tables/q3_tables.tex",ROOT/"tables/q3_moisture.csv"]
    paths += [ROOT/f"reports/q3_{name}.json" for name in reports]
    audit={"status":"passed_Q23_full_workbooks_and_Q3_numerical_event_acceptance",
        "source_fingerprint":fingerprint(),"script_sha256":digest(__file__),
        "numerical_cells_read_back":counts,"q2_data_rows_per_sheet":len(full["time_s"])-1,
        "q3_data_rows":len(np.arange(60,int(event)+1,60))+1,"original_sources_unchanged":len(original["files"]),
        "all_previous_Q1_Q2_audited_artifacts_unchanged":True,
        "critical_time_s":event,"critical_time_h":event/3600,
        "remaining_scope":["Q4 moving material domain and verification","geometry/end-face and boundary closure validity","full paper and project audit"],
        "artifacts_sha256":{str(p.relative_to(ROOT)):digest(p) for p in paths}}
    save_json(ROOT/"reports/q3_export_audit.json",audit)
    print(json.dumps(audit,ensure_ascii=False,indent=2),flush=True)


if __name__=="__main__":
    main()
