#!/usr/bin/env python3
"""Gate Q4 and export physical radii, true surface and genuinely blank cells."""
import csv
import json
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.cell import WriteOnlyCell

from scripts.question4 import read_case,fingerprint
from src.data_io import ROOT,digest,load_config,save_json,verify_sources
from validation.q4_independent_fem import fingerprint as independent_fingerprint
from validation.q4_factorial_verification import independent_fingerprint as factorial_independent_fingerprint


def json_read(path):return json.loads(Path(path).read_text())


def main():
    inputs=verify_sources();cfg=load_config();cache,meta=read_case("q4_final")
    reports={name:json_read(ROOT/f"reports/q4_{name}.json") for name in
        ("convergence","kernel_checks","independent_verification","sensitivity","factorial","factorial_verification","diagnostics")}
    for name,report in reports.items():
        script=ROOT/("scripts/factorial_q4.py" if name=="factorial" else f"validation/q4_{name}.py")
        assert report["source_fingerprint"]==fingerprint() and report["script_sha256"]==digest(script),name
        if "complete" in report:assert report["complete"],name
    assert reports["kernel_checks"]["status"].startswith("passed_")
    assert reports["diagnostics"]["status"].startswith("passed_")
    assert meta["n"]==1280 and meta["property_set"]=="q4" and meta["geometry"]=="linear"
    assert meta["radius_extension"]=="error" and meta["initialization"]=="fresh" and meta["physical_time_origin_s"]==0.
    assert meta["boundary_scenario"]=="main" and meta["perturbation"]=={}
    assert (meta["rtol"],meta["atol_temperature"],meta["atol_moisture"],meta["max_step_s"],meta["method"])==(1e-10,1e-11,1e-13,30.,"BDF")
    st=meta["statistics"];event=st["event_time_s"]
    assert st["event_found"] and event<259200. and st["water_balance_relative"]<1e-6 and st["heat_corrected_balance_relative"]<1e-6
    assert st["post_event_g"][0]>0 and abs(st["post_event_g"][1])<1e-10 and max(st["post_event_g"][2:])<0
    assert np.array_equal(cache["profile_state"][0],np.tile([28.,2.55],1280))
    labels=set()
    conv=reports["convergence"]
    for row in conv["mesh_series"]:
        labels.add(row["label"])
        if row["n"]>=1280:
            assert all(r["max_abs"]<5e-5 for r in row["difference_from_previous"]["fields"].values())
    assert abs(conv["mesh_series"][-1]["difference_from_previous"]["event_delta_s"])<.18
    for row in conv["temporal"]:
        labels.add(row["label"])
        if row["label"] in ("q4_cap60","q4_cap15","q4_radau"):
            assert abs(row["difference"]["event_delta_s"])<.18
            assert all(r["max_abs"]<5e-5 for r in row["difference"]["fields"].values())
    for name in ("sensitivity","factorial"):
        labels.update(row["label"] for row in reports[name]["cases"])
    assert len(reports["sensitivity"]["cases"])==11
    labels.update(row["fine_label"] for row in reports["factorial"]["refinements"])
    for row in reports["factorial_verification"]["cases"]:
        labels.add(row["cap15_label"])
        assert abs(row["primary_time_comparison"]["event_delta_s"])<.18
        finest=row["independent"][-1]
        assert finest["n"]==2560 and abs(finest["event_delta_primary_s"])<.18
        assert max(finest["sampled_field_max_differences"].values())<5e-5
        for item in row["independent"]:
            folder=ROOT/"results/cache/q4_factorial_independent"
            im=json_read(folder/(item["label"]+".json"))
            assert im["source_fingerprint"]==factorial_independent_fingerprint()
            assert im["cache_sha256"]==digest(folder/(item["label"]+".npz"))
    for label in labels:read_case(label)
    indep=reports["independent_verification"]
    assert indep["independent_fingerprint"]==independent_fingerprint()
    for row in indep["cases"]:
        folder=ROOT/"results/cache/q4_independent";im=json_read(folder/(row["label"]+".json"))
        assert im["source_fingerprint"]==independent_fingerprint()
        assert im["cache_sha256"]==digest(folder/(row["label"]+".npz"))
    finest=indep["cases"][-1]
    assert abs(finest["event_delta_primary_s"])<.18
    assert all(r["max_abs"]<5e-5 for r in finest["sampled_comparisons"].values())
    reproduction,rmeta=read_case("q4_reproduction")
    for key,value in cache.items():
        assert np.array_equal(value,reproduction[key],equal_nan=True),("reproduction",key)
    assert rmeta["statistics"]["event_time_s"]==event
    # No coarse-grid row and no diagnostic 1/10 s rows enter the requested workbook.
    times=cache["time_s"];R=cache["radius_current_m"]
    indices=np.flatnonzero((times>0)&((times%60==0)|(times==event)))
    expected_times=np.r_[np.arange(60.,event,60.),event]
    assert np.array_equal(times[indices],expected_times)
    valid=cache["radius_fixed_m"][None,:]<=R[:,None]
    assert np.array_equal(np.isfinite(cache["moisture"][:,:21]),valid)
    assert np.isfinite(cache["moisture"][:,-1]).all()
    template=load_workbook(Path(cfg["source_directory"])/"result4.xlsx",read_only=True)
    names=template.sheetnames;first=template.worksheets[0].cell(1,1).value
    surface_header=list(template.worksheets[0].values)[0][-1];template.close()
    assert names==["Sheet1"] and surface_header=="药材表面"
    output=ROOT/"results/final";output.mkdir(parents=True,exist_ok=True)
    wb=Workbook(write_only=True);wb.properties.creator="MathModel"
    wb.properties.description="Q4 observed radial shrinkage. Fixed-radius points outside current R(t) are blank; terminal row is critical equality."
    ws=wb.create_sheet(names[0]);ws.freeze_panes="B2"
    header=[first]+[round(j*.1,1) for j in range(21)]+[surface_header]
    ws.append(header)
    for i in indices:
        tc=WriteOnlyCell(ws,value=float(times[i]));tc.number_format="0" if times[i]==int(times[i]) else "0.0000"
        row=[tc]
        for value in cache["moisture"][i]:
            cell=WriteOnlyCell(ws,value=None if np.isnan(value) else float(np.round(value,4)))
            cell.number_format="0.0000";row.append(cell)
        ws.append(row)
    path=output/"result4.xlsx";wb.save(path);wb.close()
    # Read every populated AND blank field back; zero/formula/replicated surface is rejected.
    wb=load_workbook(path,read_only=True,data_only=False);assert wb.sheetnames==names
    rows=wb[names[0]].iter_rows();assert [c.value for c in next(rows)]==header
    numeric=blank=0
    for count,(i,row) in enumerate(zip(indices,rows),1):
        assert len(row)==23 and abs(row[0].value-times[i])<1e-8
        for j,cell in enumerate(row[1:]):
            value=cache["moisture"][i,j]
            if np.isnan(value):
                assert cell.value is None;blank+=1
            else:
                assert cell.data_type=="n" and cell.number_format=="0.0000"
                assert cell.value==float(np.round(value,4));numeric+=1
    assert count==len(indices) and next(rows,None) is None
    wb.close();assert numeric+blank==22*len(indices)
    np.savez_compressed(output/"q4_full_precision.npz",**cache)
    # Paper table uses fixed physical r, NOT material coordinate xi.
    paper_times=np.r_[np.arange(21600.,event,21600.),event]
    pi=np.searchsorted(times,paper_times);assert np.array_equal(times[pi],paper_times)
    columns=[0,5,10,15,20,21]
    values=cache["moisture"][np.ix_(pi,columns)]
    paper_rows=[[f"{time/3600:.4f}"]+["" if np.isnan(v) else f"{v:.4f}" for v in row] for time,row in zip(paper_times,values)]
    paper_header=["时间/h","0 cm","0.5 cm","1 cm","1.5 cm","2 cm","药材表面"]
    with (ROOT/"tables/q4_moisture.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f);writer.writerow(paper_header);writer.writerows(paper_rows)
    md=["# 问题四：论文表6","","固定物理距离；空白表示已在药材域外，最后一列是当前真实表面。末行为临界Cmax=0.15，不是严格低于阈值。","",
        "| "+" | ".join(paper_header)+" |","|"+"---|"*7]+["| "+" | ".join(row)+" |" for row in paper_rows]
    (ROOT/"tables/q4_tables.md").write_text("\n".join(md)+"\n",encoding="utf-8")
    tex=[r"\begin{table}[htbp]",r"\centering",r"\caption{收缩药材的干基含水率}",r"\label{tab:q4_moisture}",
         r"\begin{tabular}{rrrrrrr}",r"\toprule",r"时间/h & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm & 药材表面 \\",r"\midrule"]
    tex += [" & ".join(row)+r" \\" for row in paper_rows]
    tex += [r"\bottomrule",r"\end{tabular}",r"\end{table}",""]
    (ROOT/"tables/q4_tables.tex").write_text("\n".join(tex),encoding="utf-8")
    with (ROOT/"tables/q4_radius.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f);writer.writerow(["time_h","current_radius_cm"])
        writer.writerows([[f"{time/3600:.4f}",f"{rad*100:.6f}"] for time,rad in zip(paper_times,R[pi])])
    # Independent CSV parse verifies values and absence of substitutions.
    with (ROOT/"tables/q4_moisture.csv").open(encoding="utf-8-sig",newline="") as f:
        rows=list(csv.reader(f))[1:]
    for row,time,field in zip(rows,paper_times,values):
        assert row[0]==f"{time/3600:.4f}"
        for item,value in zip(row[1:],field):
            assert item==("" if np.isnan(value) else f"{value:.4f}")
    factor=reports["factorial"]
    with (ROOT/"tables/q4_factorial.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f);writer.writerow(["scenario","properties","geometry","critical_time_h","radius_cm"])
        writer.writerows([[r["letter"],r["group"],r["geometry"],r["event_h"],r["event_radius_cm"]] for r in factor["cases"]])
    resolved={"status":"Q4_and_factorial_numerical_milestone_accepted_physical_validity_pending",
        "source_fingerprint":fingerprint(),"parameters":{k:meta[k] for k in ("n","grading","property_set","geometry","radius_extension","rtol","atol_temperature","atol_moisture","max_step_s","method","boundary_scenario","initialization")},
        "critical_time_s":event,"critical_time_h":event/3600,"critical_radius_cm":float(R[-1]*100),
        "first_verified_strict_integer_second":st["first_strict_integer_second"],
        "radius_data_extrapolated":False,"all_export_last_digits_proven_stable":False,
        "continuous_strict_minimum_exists":False,"empirical_rho_is_effective_thermal_coefficient_not_strict_total_density":True}
    save_json(ROOT/"config/q4_final.json",resolved)
    verify_sources();old_checks=0
    for name in ("q1_export_audit.json","q2_export_audit.json","q3_export_audit.json"):
        old=json_read(ROOT/"reports"/name)
        for p,h in old["artifacts_sha256"].items():
            assert digest(ROOT/p)==h,p;old_checks+=1
    paths=[path,output/"q4_full_precision.npz",ROOT/"config/q4_final.json",ROOT/"tables/q4_tables.md",
        ROOT/"tables/q4_tables.tex",ROOT/"tables/q4_moisture.csv",ROOT/"tables/q4_radius.csv",ROOT/"tables/q4_factorial.csv",ROOT/"reports/q4_mechanisms.json"]
    paths += [ROOT/f"reports/q4_{name}.json" for name in reports]
    audit={"status":"passed_Q4_moving_domain_full_output_and_factorial_numerical_acceptance",
        "source_fingerprint":fingerprint(),"script_sha256":digest(__file__),
        "data_rows":len(indices),"fixed_radius_columns":21,"surface_columns":1,
        "numeric_cells_read_back":numeric,"outside_domain_blank_cells_read_back":blank,
        "paper_rows":len(paper_rows),"prior_audited_artifacts_unchanged":old_checks,
        "original_inputs_unchanged":len(inputs["files"]),"fresh_reproduction_identical":True,
        "critical_time_s":event,"critical_time_h":event/3600,"critical_radius_cm":float(R[-1]*100),
        "remaining_scope":["axisymmetric/end-face assumption validation","physical boundary and empirical property validity","full paper, figures, references and final audit"],
        "artifacts_sha256":{str(p.relative_to(ROOT)):digest(p) for p in paths}}
    save_json(ROOT/"reports/q4_export_audit.json",audit);print(json.dumps(audit,ensure_ascii=False,indent=2),flush=True)


if __name__=="__main__":main()
