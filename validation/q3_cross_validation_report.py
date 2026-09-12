#!/usr/bin/env python3
"""Consolidate the independent Q3 checks into one reproducible audit report."""

from __future__ import annotations

import json
import hashlib
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def read_json(relative: str) -> dict:
    return json.loads((ROOT / relative).read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> None:
    baseline = read_json("results/cache/q3/q3_final.json")
    repeat = read_json("results/cache/q3/q3_crosscheck_20260912.json")
    convergence = read_json("reports/q3_convergence.json")
    independent = read_json("reports/q3_independent_verification.json")
    event_checks = read_json("reports/q3_event_checks.json")

    base_time = float(baseline["statistics"]["event_time_s"])
    mesh2560 = next(row for row in convergence["mesh_series"] if row["n"] == 2560)
    radau = next(row for row in convergence["temporal"] if row["label"] == "q3_radau")
    fem_fine = next(
        row
        for row in independent["cases"]
        if row["elements"] == 2560 and row["max_step_s"] == 3.75
    )

    methods = [
        {
            "role": "主解",
            "scheme": "环形FVM，N=1280；BDF，最大步长30 s",
            "event_time_s": base_time,
            "delta_from_primary_s": 0.0,
        },
        {
            "role": "全新重复计算",
            "scheme": "同主解；新标签、新缓存、重新积分",
            "event_time_s": float(repeat["statistics"]["event_time_s"]),
            "delta_from_primary_s": float(repeat["statistics"]["event_time_s"]) - base_time,
        },
        {
            "role": "积分器替换",
            "scheme": "环形FVM，N=1280；Radau，最大步长30 s",
            "event_time_s": base_time + float(radau["event_delta_s"]),
            "delta_from_primary_s": float(radau["event_delta_s"]),
        },
        {
            "role": "空间加密",
            "scheme": "环形FVM，N=2560；BDF，最大步长30 s",
            "event_time_s": float(mesh2560["event_time_s"]),
            "delta_from_primary_s": float(mesh2560["event_time_s"]) - base_time,
        },
        {
            "role": "独立离散与程序",
            "scheme": "P1有限元，N=2560；自编BDF2，步长3.75 s",
            "event_time_s": float(fem_fine["event_time_s"]),
            "delta_from_primary_s": float(fem_fine["event_delta_from_primary_s"]),
        },
    ]
    for row in methods:
        row["event_time_h"] = row["event_time_s"] / 3600.0
        row["display_4dp_h"] = round(row["event_time_h"], 4)

    with np.load(ROOT / "results/final/q23_full_precision.npz") as cache:
        final_c = cache["moisture"][-1]
        final_summary = {
            "maximum": float(cache["moisture_global_max"][-1]),
            "maximum_output_radius_cm": float(cache["radius_m"][np.argmax(final_c)] * 100),
            "surface": float(final_c[-1]),
            "cross_section_mean": float(cache["moisture_mean"][-1]),
            "radial_profile_nonincreasing": bool(np.all(np.diff(final_c) <= 1e-12)),
        }

    result3 = ROOT / "results/final/result3.xlsx"
    submitted3 = ROOT / "results/submission/result3.xlsx"
    worksheet = load_workbook(result3, read_only=True, data_only=True).active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook_last = rows[-1]
    workbook = {
        "data_rows": len(rows) - 1,
        "last_time_s": float(workbook_last[0]),
        "last_row_maximum": float(max(v for v in workbook_last[1:] if isinstance(v, (int, float)))),
        "final_sha256": sha256(result3),
        "submission_sha256": sha256(submitted3),
    }

    g = [float(x) for x in baseline["statistics"]["post_event_g"]]
    checks = {
        "fresh_repeat_bitwise_time_match": methods[1]["delta_from_primary_s"] == 0.0,
        "independent_fem_within_0p01_s": abs(methods[4]["delta_from_primary_s"]) < 0.01,
        "fine_primary_mesh_within_0p1_s": abs(methods[3]["delta_from_primary_s"]) < 0.1,
        "radau_within_0p001_s": abs(methods[2]["delta_from_primary_s"]) < 0.001,
        "all_methods_round_to_57p4740_h": all(row["display_4dp_h"] == 57.4740 for row in methods),
        "event_sign_change_verified": g[0] > 0 and abs(g[1]) < 1e-12 and g[2] < 0 and g[3] < 0,
        "global_max_monotone_on_accepted_nodes_and_midpoints": bool(
            event_checks["global_max_decreased_at_all_accepted_and_midpoint_checks"]
        ),
        "manufactured_interior_peak_detected_exactly": float(
            event_checks["global_max_difference_from_PCHIP_stationary_point_search"]
        ) < 1e-12,
        "water_balance_relative_below_1e_minus_8": float(
            baseline["statistics"]["water_balance_relative"]
        ) < 1e-8,
        "terminal_profile_nonincreasing": final_summary["radial_profile_nonincreasing"],
        "workbook_terminal_time_matches_event": workbook["last_time_s"] == base_time,
        "workbook_terminal_maximum_is_0p15": workbook["last_row_maximum"] == 0.15,
        "final_and_submission_workbooks_identical": workbook["final_sha256"] == workbook["submission_sha256"],
    }
    status = "passed" if all(checks.values()) else "failed"
    report = {
        "status": status,
        "scope": "Q3 numerical and event cross-validation; not experimental validation",
        "critical_time_s": base_time,
        "critical_time_h": base_time / 3600.0,
        "first_strict_integer_second": baseline["statistics"]["first_checked_integer_second"],
        "methods": methods,
        "event_neighborhood": {
            "times_s": baseline["statistics"]["post_event_times_s"],
            "g_equals_maxC_minus_0p15": g,
            "slope_per_s": baseline["statistics"]["event_slope_per_s"],
        },
        "final_state": final_summary,
        "water_balance_relative": baseline["statistics"]["water_balance_relative"],
        "independent_sampled_field_max_difference": fem_fine["sampled_field_max_difference"],
        "result3_workbook": workbook,
        "checks": checks,
        "interpretation": (
            "The selected one-dimensional model has a reproducible downward global-threshold "
            "crossing. Agreement across implementations establishes numerical robustness only; "
            "it does not replace internal-moisture experiments or remove model-form uncertainty."
        ),
    }
    if status != "passed":
        raise AssertionError(json.dumps(checks, ensure_ascii=False, indent=2))

    (ROOT / "reports/q3_cross_validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    lines = [
        "# 问题三交叉验证报告",
        "",
        "状态：通过。该结论是既定一维模型下的数值交叉验证，不是实验验证。",
        "",
        f"主解临界时间为 {base_time:.6f} s（{base_time/3600:.9f} h），首个经回代核验严格达标的整数秒为 {baseline['statistics']['first_checked_integer_second']:.0f} s。",
        "",
        "| 证据 | 离散与求解 | 临界时间/s | 相对主解/s | 四位小时 |",
        "|---|---|---:|---:|---:|",
    ]
    for row in methods:
        lines.append(
            f"| {row['role']} | {row['scheme']} | {row['event_time_s']:.6f} | "
            f"{row['delta_from_primary_s']:+.6f} | {row['display_4dp_h']:.4f} |"
        )
    lines.extend(
        [
            "",
            f"- 临界前1 s、临界点、206907 s、临界后1 s的事件函数值依次为：{g[0]:+.9e}、{g[1]:+.9e}、{g[2]:+.9e}、{g[3]:+.9e}。",
            f"- 主网格加密到N=2560后终点改变 {methods[3]['delta_from_primary_s']:+.6f} s；按近二阶收敛，主网格采用约0.1 s作为保守数值尺度。",
            f"- 最细独立有限元与主解相差 {methods[4]['delta_from_primary_s']:+.6f} s；公共保存时刻最大温差和含水率差分别为 {fem_fine['sampled_field_max_difference']['temperature']:.3e} ℃、{fem_fine['sampled_field_max_difference']['moisture']:.3e}。该差值用于说明两条离散路线一致，不作为误差上界。",
            f"- 终点最大值位于 r={final_summary['maximum_output_radius_cm']:.1f} cm，Cmax={final_summary['maximum']:.12f}；表面值为 {final_summary['surface']:.6f}，截面平均值为 {final_summary['cross_section_mean']:.6f}。",
            f"- 水分累计守恒相对残差为 {baseline['statistics']['water_balance_relative']:.3e}。",
            "- 人造内部湿峰反例与PCHIP驻点搜索完全一致，证明算法没有把中心点预设为最大值。",
            f"- result3.xlsx共{workbook['data_rows']}行数据，末行时间与临界事件一致；最终版和提交副本SHA-256均为 `{workbook['final_sha256']}`。",
            "",
            "结论：57.4740 h 在主网格、细网格、替代积分器及独立有限元中均保持不变；当前证据支持其作为题定显示结果，但不应解释为真实工艺具有四位小数的物理精度。",
        ]
    )
    (ROOT / "reports/q3_cross_validation.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )
    print(json.dumps({"status": status, "critical_time_s": base_time, "methods": len(methods)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
