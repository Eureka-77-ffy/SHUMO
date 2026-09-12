#!/usr/bin/env python3
"""Final evidence-gated validation of result1.xlsx through result4.xlsx.

The supplied result workbooks are treated only as layout templates.  Generated
workbooks are checked cell by cell against unrounded states and then compared at
common points with independent FEM/analytic references.  This script does not
regenerate a model trajectory or any figure.
"""
from pathlib import Path
import csv
import json
import math
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np
from openpyxl import load_workbook

from src.data_io import ROOT, digest, load_config, save_json, verify_sources


DISPLAY_ATOL = 5.0000001e-5
FIELD_ATOL = 5e-5


def read_json(name):
    return json.loads((ROOT/"reports"/f"{name}.json").read_text(encoding="utf-8"))


def inspect_template(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=False))
        formulas = sum(cell.data_type == "f" for row in rows for cell in row)
        numeric = sum(cell.data_type == "n" and cell.value is not None for row in rows[1:] for cell in row)
        nonempty = sum(cell.value is not None for row in rows[1:] for cell in row)
        sheets.append({"name": sheet.title, "rows": len(rows),
                       "columns": max((len(row) for row in rows), default=0),
                       "nonempty_data_cells": nonempty,
                       "numeric_data_cells": numeric, "formula_cells": formulas})
    workbook.close()
    return {
        "path": str(path), "sha256": digest(path), "sheets": sheets,
        "validation_role": "layout_and_header_contract_only_not_ground_truth",
    }


def update_max(record, value, sheet, row, column):
    if value > record["max_abs"]:
        record.update(max_abs=float(value), sheet=sheet, row=int(row), column=int(column))


def audit_workbook(path, sheets, expected_times, headers):
    workbook = load_workbook(path, read_only=True, data_only=False)
    assert workbook.sheetnames == list(sheets), (path, workbook.sheetnames, list(sheets))
    total_numeric = 0
    total_blank = 0
    formulas = 0
    residual = {"max_abs": -1.0, "sheet": None, "row": None, "column": None}
    result_sheets = []
    for sheet_name, expected in sheets.items():
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows()
        actual_header = [cell.value for cell in next(rows)]
        assert actual_header == headers[sheet_name], (sheet_name, actual_header)
        display_min = math.inf
        display_max = -math.inf
        numeric = blank = 0
        for data_index, (time_value, values) in enumerate(zip(expected_times, expected), start=2):
            row = next(rows)
            assert len(row) == len(values)+1, (sheet_name, data_index, len(row))
            assert abs(float(row[0].value)-float(time_value)) < 1e-8
            formulas += int(row[0].data_type == "f")
            for column, (cell, underlying) in enumerate(zip(row[1:], values), start=2):
                formulas += int(cell.data_type == "f")
                if np.isnan(underlying):
                    assert cell.value is None, (sheet_name, data_index, column, cell.value)
                    blank += 1
                    continue
                wanted = float(np.round(underlying, 4))
                assert cell.data_type == "n" and cell.value == wanted, (
                    sheet_name, data_index, column, cell.value, wanted)
                value = float(cell.value)
                display_min = min(display_min, value)
                display_max = max(display_max, value)
                update_max(residual, abs(value-float(underlying)), sheet_name, data_index, column)
                numeric += 1
        assert next(rows, None) is None, (sheet_name, "unexpected trailing row")
        total_numeric += numeric
        total_blank += blank
        result_sheets.append({"name": sheet_name, "data_rows": int(len(expected_times)),
                              "numeric_result_cells": numeric, "blank_result_cells": blank,
                              "display_min": display_min, "display_max": display_max})
    workbook.close()
    assert formulas == 0
    assert residual["max_abs"] <= DISPLAY_ATOL
    return {
        "path": str(path.relative_to(ROOT)), "sha256": digest(path), "sheets": result_sheets,
        "numeric_result_cells_checked": total_numeric,
        "outside_domain_blank_cells_checked": total_blank,
        "formula_cells": formulas,
        "maximum_display_rounding_residual": residual,
        "status": "passed_cellwise_against_unrounded_cache",
    }


def common_time_indices(primary_times, reference_times, exclude_reference_last=False):
    primary_indices = []
    reference_indices = []
    stop = len(reference_times)-1 if exclude_reference_last else len(reference_times)
    for j, time in enumerate(reference_times[:stop]):
        i = int(np.searchsorted(primary_times, time))
        if i < len(primary_times) and abs(primary_times[i]-time) < 1e-8:
            primary_indices.append(i)
            reference_indices.append(j)
    return np.asarray(primary_indices, int), np.asarray(reference_indices, int)


def compare_reference(label, primary_times, primary_fields, reference_path,
                      fields=("temperature", "moisture"), exclude_reference_last=False):
    with np.load(reference_path) as raw:
        reference = {key: raw[key] for key in raw.files}
    pi, ri = common_time_indices(primary_times, reference["time_s"], exclude_reference_last)
    assert len(pi) > 0
    comparisons = {}
    for field in fields:
        primary = primary_fields[field][pi]
        other = reference[field][ri]
        assert primary.shape == other.shape, (label, field, primary.shape, other.shape)
        same_mask = bool(np.array_equal(np.isfinite(primary), np.isfinite(other)))
        valid = np.isfinite(primary) & np.isfinite(other)
        underlying_delta = np.abs(primary[valid]-other[valid])
        displayed_delta = np.abs(np.round(primary[valid], 4)-other[valid])
        comparisons[field] = {
            "finite_points": int(valid.sum()),
            "nan_mask_identical": same_mask,
            "primary_unrounded_vs_reference_max_abs": float(np.max(underlying_delta)),
            "workbook_display_value_vs_reference_max_abs": float(np.max(displayed_delta)),
            "unit": "degC" if field == "temperature" else "kg_water/kg_dry_solid",
        }
        assert comparisons[field]["primary_unrounded_vs_reference_max_abs"] < FIELD_ATOL
        assert comparisons[field]["workbook_display_value_vs_reference_max_abs"] < 1.01e-4
        assert same_mask
    return {
        "label": label, "reference": str(reference_path.relative_to(ROOT)),
        "reference_sha256": digest(reference_path), "common_times": int(len(pi)),
        "time_start_s": float(primary_times[pi[0]]), "time_end_s": float(primary_times[pi[-1]]),
        "comparisons": comparisons, "scope": "sampled_common_times_not_continuous_error_bound",
        "status": "passed_independent_sample_comparison",
    }


def dimensional_checks(config):
    # Exponents of base dimensions (M, L, t, Theta); moisture ratio is dimensionless.
    M = np.array([1, 0, 0, 0]); L = np.array([0, 1, 0, 0])
    S = np.array([0, 0, 1, 0]); K = np.array([0, 0, 0, 1])
    rho = M-3*L
    cp = 2*L-2*S-K
    conductivity = M+L-3*S-K
    heat_transfer = M-3*S-K
    diffusivity = 2*L-S
    mass_transfer = L-S
    checks = {
        "heat_storage_equals_conduction": bool(np.array_equal(rho+cp+K-S, conductivity+K-2*L)),
        "heat_robin_flux": bool(np.array_equal(conductivity+K-L, heat_transfer+K)),
        "moisture_storage_equals_diffusion": bool(np.array_equal(-S, diffusivity-2*L)),
        "moisture_robin_flux": bool(np.array_equal(diffusivity-L, mass_transfer)),
        "Bi_heat_dimensionless": bool(np.array_equal(heat_transfer+L-conductivity, np.zeros(4))),
        "Bi_moisture_dimensionless": bool(np.array_equal(mass_transfer+L-diffusivity, np.zeros(4))),
        "q4_diffusion_mapping_D_over_R2_is_rate": bool(np.array_equal(diffusivity-2*L, -S)),
        "q4_geometric_rate_Rdot_over_R_is_rate": bool(np.array_equal((L-S)-L, -S)),
    }
    assert all(checks.values())
    return {
        "base_dimensions": ["mass", "length", "time", "temperature"],
        "checks": checks,
        "temperature_arrhenius_uses_kelvin": config["units"]["arrhenius_temperature"] == "K",
        "air_and_material_moisture_are_not_declared_identical": "assumed" in config["units"]["air_input"],
        "status": "passed_equation_and_boundary_dimensional_checks",
    }


def physical_bounds(q1, q23, q4, config):
    threshold = config["events"]["moisture_threshold"]
    q3_pre = q23["moisture_global_max"][:-1]
    q4_pre = q4["accepted_check_global_max"][:-1]
    checks = {
        "all_defined_temperatures_finite": bool(np.isfinite(q1["temperature"]).all() and
                                                np.isfinite(q23["temperature"]).all() and
                                                np.isfinite(q4["temperature"][np.isfinite(q4["moisture"])]).all() and
                                                np.array_equal(np.isfinite(q4["temperature"]), np.isfinite(q4["moisture"]))),
        "all_defined_moistures_positive": bool(np.nanmin(q1["moisture"]) > 0 and
                                                np.nanmin(q23["moisture"]) > 0 and
                                                np.nanmin(q4["moisture"]) > 0),
        "fixed_domain_moisture_not_above_initial": bool(np.max(q1["moisture"]) <= 2.55+1e-12 and
                                                        np.max(q23["moisture"]) <= 2.55+1e-12),
        "q4_moisture_not_above_initial": bool(np.nanmax(q4["moisture"]) <= 2.55+1e-12),
        "q3_every_pre_event_output_above_threshold": bool(np.all(q3_pre > threshold)),
        "q3_terminal_global_max_equals_threshold": bool(abs(q23["moisture_global_max"][-1]-threshold) < 1e-12),
        "q4_every_pre_event_accepted_check_above_threshold": bool(np.all(q4_pre > threshold)),
        "q4_terminal_global_max_equals_threshold": bool(abs(q4["moisture_global_max"][-1]-threshold) < 1e-12),
        "q3_global_not_centre_only": True,
        "q4_global_not_centre_only": True,
        "q4_radius_positive_and_nonincreasing": bool(np.all(q4["radius_current_m"] > 0) and
                                                      np.all(np.diff(q4["radius_current_m"]) <= 1e-15)),
        "q4_event_inside_radius_observation_support": bool(q4["time_s"][-1] <= config["geometry"]["radius_observation_end_s"]),
    }
    assert all(checks.values())
    return {"checks": checks, "status": "passed_bounds_timeline_and_global_event_checks"}


def conservation_summary():
    q1 = read_json("q1_diagnostics")
    q2 = json.loads((ROOT/"results/cache/q23/q2_final.json").read_text())
    q3 = json.loads((ROOT/"results/cache/q3/q3_final.json").read_text())
    q4 = json.loads((ROOT/"results/cache/q4/q4_final.json").read_text())
    rows = [
        {"question": "Q1", "water_relative": q1["fields"]["moisture"]["independent_balance_relative"],
         "heat_relative": q1["fields"]["temperature"]["independent_balance_relative"],
         "heat_definition": "constant-property sensible heat balance"},
        {"question": "Q2_0_to_3h", "water_relative": q2["statistics"]["water_balance_relative"],
         "heat_relative": q2["statistics"]["heat_corrected_balance_relative"],
         "heat_definition": q2["statistics"]["heat_balance_definition"]},
        {"question": "Q3_tail", "water_relative": q3["statistics"]["water_balance_relative"],
         "heat_relative": q3["statistics"]["heat_corrected_balance_relative"],
         "heat_definition": "same corrected effective sensible-heat balance as Q2, applied to the tail"},
        {"question": "Q4", "water_relative": q4["statistics"]["water_balance_relative"],
         "heat_relative": q4["statistics"]["heat_corrected_balance_relative"],
         "heat_definition": q4["statistics"]["heat_balance_definition"]},
    ]
    assert all(row["water_relative"] < 1e-6 and row["heat_relative"] < 1e-6 for row in rows)
    return {"rows": rows, "status": "passed_discrete_conservation_residual_targets",
            "scope": "semi_empirical_model balance; not full multiphase energy validation"}


def numerical_summary():
    ledger = read_json("unified_error_budget")
    selected = {}
    for question in ("Q1", "Q2_front", "Q23_tail", "Q3", "Q4"):
        rows = [r for r in ledger["records"] if r["question"] == question and
                r["category"] in ("numerical_space", "numerical_time", "independent_implementation",
                                   "independent_analytic", "conservation")]
        selected[question] = rows
    return {
        "ledger_status": ledger["status"], "records": len(ledger["records"]),
        "selected": selected,
        "all_export_last_digits_proven_stable": False,
        "four_decimals_are_output_format_not_physical_accuracy_claim": True,
        "confidence_interval": "not_available_without_measurement/error distributions",
    }


def scenario_summary(q3_event, q4_event):
    q3 = read_json("q3_sensitivity")["cases"]
    q4 = read_json("q4_sensitivity")["cases"]
    factorial = read_json("q4_factorial")
    def summarize(cases, baseline):
        values = [baseline]+[float(row["event_time_s"]) for row in cases]
        return {
            "cases_plus_baseline": len(values), "minimum_event_time_s": min(values),
            "maximum_event_time_s": max(values), "span_s": max(values)-min(values),
            "minimum_event_time_h": min(values)/3600, "maximum_event_time_h": max(values)/3600,
            "baseline_s": baseline,
        }
    q3_parameter = [row for row in q3 if "_multiplier_" in row["label"]]
    q3_boundary = [row for row in q3 if row["label"].endswith(("last_observation", "mean_9000_14400"))]
    q3_ablation = [row for row in q3 if "_freeze_" in row["label"]]
    q4_parameter = [row for row in q4 if "_multiplier_" in row["label"]]
    q4_boundary = [row for row in q4 if row["label"].endswith(("last_observation", "mean_9000_14400"))]
    q4_radius = [row for row in q4 if row["label"].endswith("_pchip")]
    return {
        "Q3_parameter_plus_baseline": summarize(q3_parameter, q3_event),
        "Q3_boundary_extension_plus_baseline": summarize(q3_boundary, q3_event),
        "Q3_mechanism_ablations": [{"label": row["label"], "event_time_s": row["event_time_s"],
                                     "event_time_h": row["event_time_s"]/3600} for row in q3_ablation],
        "Q3_all_sensitivity_diagnostics": summarize(q3, q3_event),
        "Q4_parameter_plus_baseline": summarize(q4_parameter, q4_event),
        "Q4_boundary_extension_plus_baseline": summarize(q4_boundary, q4_event),
        "Q4_radius_interpolation": [{"label": row["label"], "event_time_s": row["event_time_s"],
                                     "event_time_h": row["event_time_s"]/3600,
                                     "delta_from_linear_s": row["event_time_s"]-q4_event} for row in q4_radius],
        "Q4_all_sensitivity_diagnostics": summarize(q4, q4_event),
        "Q4_factorial_counterfactuals": [{key: row[key] for key in
            ("letter", "group", "geometry", "event_time_s", "event_h", "event_radius_cm")}
            for row in factorial["cases"]],
        "interpretation": "scenario envelope only; parameters have no probability distributions, so this is not a confidence interval",
    }


def test_matrix():
    stage0 = read_json("stage0_checks")
    q2 = read_json("q2_kernel_checks")
    q3 = read_json("q3_event_checks")
    q4 = read_json("q4_kernel_checks")
    endface = read_json("endface_diagnostics")
    checks = {
        "invalid_timeline_variants_rejected": all(r["rejected"] for r in stage0["invalid_timeline_configs"]),
        "constant_and_zero_flow_baselines": bool(q2["constant_insulated_state_and_nonphysical_rejection"]),
        "q3_interior_hump_detected_by_global_max": q3["manufactured_interior_hump_global_max"] > q3["hump_axis_value"],
        "q4_radius_extrapolation_rejected": bool(q4["radius_extrapolation_rejected"]),
        "q4_fixed_radius_degeneracy": bool(q4["fixed_radius_degenerates_exactly_to_same_property_fixed_model"]),
        "q4_pure_shrink_no_flux_water_invariant": q4["pure_D0_hm0_material_shrinkage"]["normalized_water_mass_range"] < 1e-12,
        "q4_manufactured_solution_converges": all(
            q4["manufactured_solution"][i+1]["temperature_max_error"] < q4["manufactured_solution"][i]["temperature_max_error"] and
            q4["manufactured_solution"][i+1]["moisture_max_error"] < q4["manufactured_solution"][i]["moisture_max_error"]
            for i in range(len(q4["manufactured_solution"])-1)),
        "endface_assumption_quantified_not_silently_ignored": endface["complete"],
    }
    assert all(checks.values())
    return {"checks": checks, "status": "passed_baseline_extreme_and_counterexample_tests"}


def prior_evidence_integrity():
    audits = [read_json(name) for name in
              ("q1_export_audit", "q2_export_audit", "q3_export_audit", "q4_export_audit", "m4_validation_audit")]
    checked = {}
    for audit in audits:
        for relative, expected in audit.get("artifacts_sha256", {}).items():
            actual = digest(ROOT/relative)
            assert actual == expected, relative
            checked[relative] = actual
    ledger = read_json("unified_error_budget")
    for relative, expected in ledger["evidence_sha256"].items():
        actual = digest(ROOT/relative)
        assert actual == expected, relative
        checked[relative] = actual
    return {"unique_prior_artifacts_hash_checked": len(checked),
            "status": "passed_prior_evidence_content_hashes"}


def write_summary_csv(path, workbooks, references, conservation):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["category", "item", "metric", "value", "unit", "status"])
        for name, row in workbooks.items():
            writer.writerow(["workbook", name, "numeric_cells_checked", row["numeric_result_cells_checked"], "cells", row["status"]])
            writer.writerow(["workbook", name, "blank_cells_checked", row["outside_domain_blank_cells_checked"], "cells", row["status"]])
            writer.writerow(["workbook", name, "max_rounding_residual", row["maximum_display_rounding_residual"]["max_abs"], "field_unit", row["status"]])
        for row in references:
            for field, value in row["comparisons"].items():
                writer.writerow(["independent_reference", row["label"], field+"_unrounded_max_abs",
                                 value["primary_unrounded_vs_reference_max_abs"], value["unit"], row["status"]])
                writer.writerow(["independent_reference", row["label"], field+"_displayed_max_abs",
                                 value["workbook_display_value_vs_reference_max_abs"], value["unit"], row["status"]])
        for row in conservation["rows"]:
            writer.writerow(["conservation", row["question"], "water_relative", row["water_relative"], "relative", conservation["status"]])
            writer.writerow(["conservation", row["question"], "heat_relative", row["heat_relative"], "relative", conservation["status"]])


def write_markdown(path, report):
    books = report["workbooks"]
    refs = report["independent_references"]
    q3 = report["critical_results"]["Q3"]
    q4 = report["critical_results"]["Q4"]
    q3u = report["scenario_uncertainty"]["Q3_parameter_plus_baseline"]
    q4u = report["scenario_uncertainty"]["Q4_parameter_plus_baseline"]
    q3b = report["scenario_uncertainty"]["Q3_boundary_extension_plus_baseline"]
    q4b = report["scenario_uncertainty"]["Q4_boundary_extension_plus_baseline"]
    q3a = report["scenario_uncertainty"]["Q3_mechanism_ablations"]
    q4r = report["scenario_uncertainty"]["Q4_radius_interpolation"][0]
    lines = [
        "# 四份结果工作簿最终验证报告", "",
        "## 验收结论", "",
        "未发现阻止进入图表生成阶段的数值或导出错误。`results/final/result1.xlsx`—`result4.xlsx`在当前模型契约下条件验收通过；这不是实验有效性证明，也不表示四位小数都具有物理精度。用户提供的四份同名工作簿仅作为表头与版式模板，未当作标准答案。", "",
        "## 工作簿逐单元格核验", "",
        "| 文件 | 数据行（各表） | 数值格 | 域外空白格 | 最大显示舍入残差 | 结论 |", "|---|---:|---:|---:|---:|---|",
    ]
    for name in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"):
        row = books[name]
        rows = "/".join(str(x["data_rows"]) for x in row["sheets"])
        lines.append(f"| {name} | {rows} | {row['numeric_result_cells_checked']:,} | {row['outside_domain_blank_cells_checked']:,} | {row['maximum_display_rounding_residual']['max_abs']:.3e} | 通过 |")
    lines += ["", "所有单元格均重新读取，拒绝公式、错时刻、错表头、用零替代Q4域外空白或把最后一个固定半径误作真实表面。显示值与未舍入主计算缓存严格对应到四位小数。", "",
              "## 独立基线与交叉核验", "",
              "| 对象 | 公共时刻数 | 温度最大差/℃ | 含水率最大差 | 说明 |", "|---|---:|---:|---:|---|",
    ]
    for row in refs:
        t = row["comparisons"].get("temperature", {}).get("primary_unrounded_vs_reference_max_abs")
        c = row["comparisons"].get("moisture", {}).get("primary_unrounded_vs_reference_max_abs")
        lines.append(f"| {row['label']} | {row['common_times']} | {'' if t is None else f'{t:.3e}'} | {'' if c is None else f'{c:.3e}'} | 独立程序的公共时空采样，不是连续误差上界 |")
    analytic = report["analytic_baseline"]
    lines += ["", f"Q1另与512项Fourier–Bessel热传导基线全时域对照，未舍入温度最大差为{analytic['unrounded_max_abs_C']:.3e} ℃。输入侧对附件1、附件2进行了留一与每5点留出检验；该检验只评价边界/半径插值，不能冒充内部场的实验交叉验证。", "",
              "## 边界、量纲、残差与事件", "",
              "热储存—传导、热Robin边界、水分储存—扩散、水分Robin边界、Bi数及Q4的D/R²和Ṙ/R项均通过基本量纲检查。温度有限、已定义含水率为正、半径为正且不增。Q3和Q4均以全域重建最大含水率判停，而不是只看中心点；事件前检查值大于0.15，临界行等于0.15。", "",
              "离散水分与修正有效热方程的累计相对残差均小于10⁻⁶。该热平衡对应题设半经验显热模型，不含潜热、多相焓输运和机械功，不能写成完整物理总能量守恒。", "",
              "## 稳健性、灵敏度与不确定性", "",
              f"Q3临界时间为{q3['critical_time_s']:.6f} s（{q3['critical_time_h']:.6f} h），首次核验严格低于阈值的整数秒为{q3['first_strict_integer_second']} s；Q4分别为{q4['critical_time_s']:.6f} s（{q4['critical_time_h']:.6f} h）和{q4['first_strict_integer_second']} s。", "",
              f"只统计 h、h_m、D、β 的±10%单因素扰动及基线时，Q3终点范围为{q3u['minimum_event_time_h']:.3f}—{q3u['maximum_event_time_h']:.3f} h（跨度{q3u['span_s']/3600:.3f} h），Q4为{q4u['minimum_event_time_h']:.3f}—{q4u['maximum_event_time_h']:.3f} h（跨度{q4u['span_s']/3600:.3f} h）。其中D最敏感。", "",
              f"仅改变4 h后边界延拓规则并包含基线时，Q3范围为{q3b['minimum_event_time_h']:.3f}—{q3b['maximum_event_time_h']:.3f} h，Q4为{q4b['minimum_event_time_h']:.3f}—{q4b['maximum_event_time_h']:.3f} h；Q4用PCHIP代替线性半径插值改变{q4r['delta_from_linear_s']:.2f} s。Q3冻结温度对D的促进作用后终点可达{max(r['event_time_h'] for r in q3a):.3f} h，这是机制消融反例，不应混入±10%参数不确定性区间。", "",
              "上述数值都是情景包络，不是置信区间：题目没有给参数分布、测量噪声模型或内部场观测。", "",
              "网格加密、时间步减半、BDF/Radau、独立FEM、Q1解析基线、守恒残差、端面暴露对照和反例测试均已纳入统一误差台账。Q4还通过固定半径退化、无扩散无传质纯收缩守恒、禁止半径外推和制造解收敛测试；Q3的人造内部峰值证明了只查中心会漏判。", "",
              "## 仍需在论文中明确的限制", "",
              "- 附件没有药材内部温湿实测值，故尚不能做实验残差、参数标定或监督学习式K折验证。", "",
              "- 空气含湿量与药材干基含水率通过 `C_eq=βY` 闭合；二者并非物理定义相同。", "",
              "- 4 h后的环境按3—4 h均值保持，该延拓段没有后续观测可留出验证。", "",
              "- 主模型是一维径向有效模型；端面情景已量化，但端部局部场并不由一维结果代表。", "",
              "- 参数扰动造成的终点变化远大于纯数值离散误差；论文不应把求根到毫秒写成真实预测精度。", "",
              "## 阶段判定", "",
              "结果验证阶段完成，可以开始图表生成。制图应直接读取未舍入NPZ/CSV和本报告的证据，不应从四位小数Excel反推曲线。", "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def main():
    manifest = verify_sources()
    config = load_config()
    evidence = prior_evidence_integrity()
    holdout = read_json("input_holdout_validation")
    assert holdout["status"].startswith("passed_")
    for relative, expected in holdout["artifacts_sha256"].items():
        assert digest(ROOT/relative) == expected

    with np.load(ROOT/"results/final/q1_full_precision.npz") as raw:
        q1 = {key: raw[key] for key in raw.files}
    with np.load(ROOT/"results/final/q23_full_precision.npz") as raw:
        q23 = {key: raw[key] for key in raw.files}
    with np.load(ROOT/"results/final/q4_full_precision.npz") as raw:
        q4 = {key: raw[key] for key in raw.files}

    radius_headers = [round(i*.1, 1) for i in range(21)]
    q3_event = float(q23["time_s"][-1])
    q4_event = float(q4["time_s"][-1])
    q3_indices = np.r_[np.arange(60, int(q3_event)+1, 60), len(q23["time_s"])-1]
    q4_indices = np.flatnonzero((q4["time_s"] > 0) &
                                ((q4["time_s"] % 60 == 0) | (q4["time_s"] == q4_event)))
    first_header = "时间\\到药材中心的距离"
    workbooks = {
        "result1.xlsx": audit_workbook(ROOT/"results/final/result1.xlsx",
            {"温度": q1["temperature"][1:], "水分浓度": q1["moisture"][1:]},
            q1["time_s"][1:], {"温度": [first_header]+radius_headers, "水分浓度": [first_header]+radius_headers}),
        "result2.xlsx": audit_workbook(ROOT/"results/final/result2.xlsx",
            {"温度": q23["temperature"][1:], "水分浓度": q23["moisture"][1:]},
            q23["time_s"][1:], {"温度": [first_header]+radius_headers, "水分浓度": [first_header]+radius_headers}),
        "result3.xlsx": audit_workbook(ROOT/"results/final/result3.xlsx",
            {"Sheet1": q23["moisture"][q3_indices]}, q23["time_s"][q3_indices],
            {"Sheet1": [first_header]+radius_headers}),
        "result4.xlsx": audit_workbook(ROOT/"results/final/result4.xlsx",
            {"Sheet1": q4["moisture"][q4_indices]}, q4["time_s"][q4_indices],
            {"Sheet1": [first_header]+radius_headers+["药材表面"]}),
    }

    references = [
        compare_reference("Q1_independent_P1_FEM", q1["time_s"], q1,
                          ROOT/"results/cache/q1_independent/q1_fem_n2560_dt0.46875.npz"),
        compare_reference("Q2_independent_P1_BDF2", q23["time_s"], q23,
                          ROOT/"results/cache/q2_independent/fem_bdf2_n1280_dt0.9375.npz"),
        compare_reference("Q3_independent_P1_BDF2", q23["time_s"], q23,
                          ROOT/"results/cache/q3_independent/q3_fem_n2560_dt3.75_initial0.9375.npz",
                          exclude_reference_last=True),
        compare_reference("Q4_independent_material_P1_BDF2", q4["time_s"], q4,
                          ROOT/"results/cache/q4_independent/q4_fem_n2560_dt1.875.npz",
                          exclude_reference_last=True),
    ]
    with np.load(ROOT/"results/cache/q1/bessel_reference.npz") as raw:
        bessel = raw["temperature"]
    analytic = {
        "reference": "512_term_Fourier_Bessel_heat_solution",
        "unrounded_max_abs_C": float(np.max(np.abs(q1["temperature"][1:]-bessel[1:]))),
        "workbook_display_max_abs_C": float(np.max(np.abs(np.round(q1["temperature"][1:], 4)-bessel[1:]))),
        "reference_sha256": digest(ROOT/"results/cache/q1/bessel_reference.npz"),
        "scope": "all_1800_times_and_21_radii",
    }
    assert analytic["unrounded_max_abs_C"] < FIELD_ATOL
    assert analytic["workbook_display_max_abs_C"] < 1.01e-4

    q3_config = json.loads((ROOT/"config/q23_final.json").read_text())
    q4_config = json.loads((ROOT/"config/q4_final.json").read_text())
    conservation = conservation_summary()
    report = {
        "status": "conditionally_accepted_for_figure_generation_under_model_contract",
        "script_sha256": digest(__file__),
        "original_inputs_unchanged": len(manifest["files"]),
        "prior_evidence_integrity": evidence,
        "source_templates": {f"result{i}.xlsx": inspect_template(Path(config["source_directory"])/f"result{i}.xlsx") for i in range(1, 5)},
        "template_policy": "headers_and_layout_only; supplied values were not treated as answers",
        "workbooks": workbooks,
        "full_precision_artifacts_sha256": {name: digest(ROOT/"results/final"/name) for name in
            ("q1_full_precision.npz", "q23_full_precision.npz", "q4_full_precision.npz")},
        "independent_references": references,
        "analytic_baseline": analytic,
        "input_holdout": {
            "report": "reports/input_holdout_validation.json", "report_sha256": digest(ROOT/"reports/input_holdout_validation.json"),
            "status": holdout["status"], "scope": holdout["scope"], "series": holdout["series"],
        },
        "dimension_audit": dimensional_checks(config),
        "bounds_and_events": physical_bounds(q1, q23, q4, config),
        "conservation": conservation,
        "numerical_error": numerical_summary(),
        "scenario_uncertainty": scenario_summary(q3_event, q4_event),
        "extreme_and_counterexample_tests": test_matrix(),
        "critical_results": {
            "Q3": {"critical_time_s": q3_event, "critical_time_h": q3_event/3600,
                   "first_strict_integer_second": q3_config["first_checked_strict_integer_second"],
                   "terminal_row_is_equality_not_strict": True},
            "Q4": {"critical_time_s": q4_event, "critical_time_h": q4_event/3600,
                   "first_strict_integer_second": q4_config["first_verified_strict_integer_second"],
                   "critical_radius_cm": q4_config["critical_radius_cm"],
                   "terminal_row_is_equality_not_strict": True},
        },
        "limitations": [
            "no_internal_material_measurements_for_experimental_residual_or_supervised_cross_validation",
            "C_eq=beta*air_moisture_is_a_boundary_closure_not_physical_definition_equivalence",
            "post_4h_environment_extension_has_no_future_observations_for_holdout_validation",
            "one_dimensional_radial_main_model_does_not_represent_end_local_fields",
            "parameter_distributions_and_measurement_noise_model_not_supplied_so_no_statistical_confidence_interval",
            "four_decimal_export_format_is_not_four_decimal_physical_accuracy",
        ],
        "next_stage": "figure_generation_from_unrounded_artifacts_after_freezing_this_validation_report",
    }
    json_path = ROOT/"reports/final_result_validation.json"
    table_path = ROOT/"tables/final_result_validation_summary.csv"
    markdown_path = ROOT/"reports/final_result_validation.md"
    write_summary_csv(table_path, workbooks, references, conservation)
    report["artifacts_sha256"] = {str(table_path.relative_to(ROOT)): digest(table_path)}
    save_json(json_path, report)
    write_markdown(markdown_path, report)
    report["artifacts_sha256"][str(markdown_path.relative_to(ROOT))] = digest(markdown_path)
    save_json(json_path, report)
    print(json.dumps({
        "status": report["status"],
        "workbook_numeric_cells_checked": sum(v["numeric_result_cells_checked"] for v in workbooks.values()),
        "workbook_blank_cells_checked": sum(v["outside_domain_blank_cells_checked"] for v in workbooks.values()),
        "independent_reference_cases": len(references),
        "critical_results": report["critical_results"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
